# -*- coding: utf-8 -*-
"""
Reconstrucao horaria de vazao por modelo estrutural em espaco de estados.

VERSAO FINAL V4-CLUSTER-CORRIGIDO — 2026-08-11

Correcao motivada pelos arquivos KALMAN_OPTIMIZATION_FAILURES.csv e
KALMAN_MODEL_SELECTION_FAILURES.csv: a convergencia nativa do otimizador foi
preservada, mas componentes de variancia no limite passaram a acionar modelos
aninhados por warm start, e o condicionamento de aceitacao passou a ser medido
na matriz de correlacao dos parametros (livre de escala). O numero de condicao
bruto da covariancia OPG continua sendo salvo para auditoria. Nesta V4, a
reproducao multistart usa tolerancias combinadas absoluta e relativa, evitando
que coeficientes validamente proximos de zero sejam considerados diferentes
por uma divisao instavel por valores quase nulos.

Escopo deliberado:
  - NAO executa controle de qualidade;
  - NAO executa validacao Q50;
  - recebe a mascara de lacunas pronta na coluna de vazao (NaN);
  - ajusta os parametros exclusivamente em 2022-2023;
  - compara estruturas parcimoniosas e aceita somente solucao reproduzida por
    multiplas inicializacoes, com gradiente e covariancia auditados;
  - usa Kalman smoother para lacunas retrospectivas de 2022-2023;
  - usa previsao causal de um passo do Kalman filter para lacunas de 2024;
  - preserva integralmente todas as vazoes originalmente observadas;
  - recria de forma deterministica apenas tres covariaveis antecedentes de
    chuva de curto prazo que ainda nao existem no arquivo pre-Kalman;
  - NAO usa precipitacao acumulada em 20 dias nem sua interacao, evitando
    introduzir a memoria hidrologica de longo prazo descartada neste estudo;
  - salva serie, flags, intervalos, parametros, Q, R, inicializacao,
    convergencia, lacunas, cinco validacoes mascaradas bem-sucedidas,
    sensibilidade aos limites, diagnosticos,
    versoes de software e um relatorio-base para o material suplementar.

O arquivo de entrada esperado e:
  /content/drive/MyDrive/Colab Notebooks/Dados (1100+ NaN).xlsx
"""

# =============================================================================
# 0. INSTALACAO E GOOGLE DRIVE
# =============================================================================

import sys
import subprocess
import os

subprocess.run(
    [
        sys.executable,
        "-m",
        "pip",
        "install",
        "-q",
        "--upgrade",
        "statsmodels==0.14.6",
        "openpyxl==3.1.5",
    ],
    check=True,
)

try:
    from google.colab import drive
    IN_COLAB = True
except ImportError:
    drive = None
    IN_COLAB = False

if IN_COLAB:
    drive.mount("/content/drive")


# =============================================================================
# 1. IMPORTACOES
# =============================================================================

import hashlib
import json
import platform
import shutil
import warnings
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import scipy
import statsmodels
from scipy import stats
from statsmodels.stats.diagnostic import acorr_ljungbox
from statsmodels.tsa.statespace.structural import UnobservedComponents


# =============================================================================
# 2. CONFIGURACAO CONGELADA
# =============================================================================

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_FILE = (
    Path("/content/drive/MyDrive/Colab Notebooks/Dados (1100+ NaN).xlsx")
    if IN_COLAB
    else REPOSITORY_ROOT / "data" / "input" / "data_with_missing_flow.xlsx"
)
INPUT_FILE = Path(os.environ.get("KALMAN_INPUT_PATH", str(DEFAULT_INPUT_FILE)))
INPUT_SHEET = 0

DEFAULT_OUTPUT_DIR = (
    Path("/content/drive/MyDrive/Colab Notebooks/Kalman_Water_Research")
    if IN_COLAB
    else REPOSITORY_ROOT / "outputs" / "kalman_reconstruction"
)
OUTPUT_DIR = Path(os.environ.get("KALMAN_OUTPUT_DIR", str(DEFAULT_OUTPUT_DIR)))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_XLSX = OUTPUT_DIR / "KALMAN_RECONSTRUCAO_FINAL.xlsx"
OUTPUT_JSON = OUTPUT_DIR / "KALMAN_METADATA.json"
OUTPUT_REQUIREMENTS = OUTPUT_DIR / "KALMAN_REQUIREMENTS.txt"
OUTPUT_SUMMARY = OUTPUT_DIR / "KALMAN_MODEL_SUMMARY.txt"
OUTPUT_MODEL = OUTPUT_DIR / "KALMAN_RESULTS_2022_2023.pkl"
OUTPUT_CODE_COPY = OUTPUT_DIR / "KALMAN_WATER_RESEARCH_COLAB_EXECUTED.py"
OUTPUT_SUPPLEMENT = OUTPUT_DIR / "KALMAN_SUPPLEMENTARY_REPORT.txt"
OUTPUT_MANIFEST = OUTPUT_DIR / "KALMAN_RUN_MANIFEST.json"
OUTPUT_PACKAGE = OUTPUT_DIR / "KALMAN_SUBMISSION_PACKAGE.zip"

DATETIME_COLUMN = "datetime"
FLOW_COLUMN = "Vazão"

# Lista explicita para impedir a inclusao acidental de IDs, flags de QC ou
# vazoes defasadas. As tres covariaveis hidrologicas derivadas sao calculadas
# neste proprio script, depois da ordenacao e validacao da grade horaria.
EXOG_COLUMNS = [
    "Precipitação",
    "Temperatura Instantanea",
    "Temperatura Média",
    "Umidade Instantanea",
    "Umidade Media",
    "Sensação Termica (°F)",
    "hora_sin",
    "hora_cos",
    "Dia da Semana_quarta-feira",
    "Dia da Semana_quinta-feira",
    "Dia da Semana_segunda-feira",
    "Dia da Semana_sexta-feira",
    "Dia da Semana_sábado",
    "Dia da Semana_terça-feira",
    "Classe do dia_Feriado",
    "interacao_classe_periodo_Dia comum_Manha",
    "interacao_classe_periodo_Dia comum_Noite",
    "interacao_classe_periodo_Dia comum_Tarde",
    "interacao_classe_periodo_Feriado_Manha",
    "interacao_classe_periodo_Feriado_Noite",
    "interacao_classe_periodo_Feriado_Tarde",
    "tempo_sem_chuva",
    "precip_lag_1h",
    "precip_lag_2h",
]

ENGINEERED_EXOG_COLUMNS = [
    "tempo_sem_chuva",
    "precip_lag_1h",
    "precip_lag_2h",
]

RAW_EXOG_COLUMNS = [
    column for column in EXOG_COLUMNS if column not in ENGINEERED_EXOG_COLUMNS
]

# Duas colunas presentes na planilha nao entram no ajuste porque sao
# combinações lineares exatas de outras dummies ja incluidas:
#   Meio de semana = segunda + terca + quarta + quinta + sexta;
#   Feriado_Madrugada = Feriado - Feriado_Manha - Feriado_Noite - Feriado_Tarde.
# A exclusao a priori fornece uma matriz de desenho identificavel, mantendo
# domingo, dia comum e madrugada como categorias de referencia apropriadas.
EXCLUDED_EXACTLY_REDUNDANT_COLUMNS = [
    "Classe Semanal_Meio de semana",
    "interacao_classe_periodo_Feriado_Madrugada",
]

FEATURE_ENGINEERING_DEFINITIONS = {
    "precip_lag_1h": "Precipitacao deslocada em 1 h; borda inicial preenchida com 0",
    "precip_lag_2h": "Precipitacao deslocada em 2 h; duas bordas iniciais preenchidas com 0",
    "tempo_sem_chuva": "Numero de horas consecutivas com precipitacao igual a 0; reinicia em 0 quando P_t > 0",
}

# Auditoria congelada do arquivo oficial. Se o arquivo mudar, o script para
# antes do ajuste, evitando que resultados de bases diferentes sejam misturados.
EXPECTED_INPUT_SHA256 = (
    "78ae23aa61d19c58159d96c515da95662acc78d938655ab7ec1140ae9e172c57"
)
EXPECTED_ROWS = 26304
EXPECTED_MISSING_TOTAL = 1334
EXPECTED_MISSING_TRAIN = 825
EXPECTED_MISSING_TEST = 509
STRICT_INPUT_AUDIT = True

TRAIN_START = pd.Timestamp("2022-01-01 00:00:00")
TRAIN_END = pd.Timestamp("2023-12-31 23:00:00")
TEST_START = pd.Timestamp("2024-01-01 00:00:00")
TEST_END = pd.Timestamp("2024-12-31 23:00:00")

# Restricoes tecnicas/operacionais previamente definidas no manuscrito.
FLOW_LOWER_BOUND = 600.0
FLOW_UPPER_BOUND = None  # None = maior vazao originalmente valida do arquivo.

# Modelos candidatos parcimoniosos. A execucao diagnostica anterior mostrou
# que a tendencia deterministica nao melhorou o BIC e que sigma2.irregular e/ou
# sigma2.ar atingiam o limite zero. Por isso, o modelo completo e mantido como
# modelo-pai e seus tres modelos aninhados de fronteira sao ajustados de forma
# explicita. Solucoes de fronteira nunca sao aceitas como resultado final: elas
# servem apenas para construir warm starts por nome de parametro para o modelo
# reduzido correspondente.
USE_LOG_FLOW = True
USE_EXACT_DIFFUSE = True
CANDIDATE_MODELS = [
    {
        "model_id": "local_level_ar1_irregular",
        "trend": False,
        "stochastic_trend": False,
        "irregular": True,
        "ar_order": 1,
        "warm_start_source": None,
        "warm_start_boundary_parameters": [],
    },
    {
        "model_id": "local_level_ar1_R_fixed_zero",
        "trend": False,
        "stochastic_trend": False,
        "irregular": False,
        "ar_order": 1,
        "warm_start_source": "local_level_ar1_irregular",
        "warm_start_boundary_parameters": ["sigma2.irregular"],
    },
    {
        "model_id": "local_level_irregular_no_ar",
        "trend": False,
        "stochastic_trend": False,
        "irregular": True,
        "ar_order": 0,
        "warm_start_source": "local_level_ar1_irregular",
        "warm_start_boundary_parameters": ["sigma2.ar"],
    },
    {
        "model_id": "local_level_R_fixed_zero_no_ar",
        "trend": False,
        "stochastic_trend": False,
        "irregular": False,
        "ar_order": 0,
        "warm_start_source": "local_level_ar1_irregular",
        "warm_start_boundary_parameters": [
            "sigma2.irregular",
            "sigma2.ar",
        ],
    },
]

# Otimizacao.
OPTIMIZER = "lbfgs"
OPTIM_SCORE_METHOD = None
MAXITER_FINAL = 3000
N_STARTS_PER_CANDIDATE = 5
MIN_STABLE_STARTS = 3
RANDOM_SEED = 20260809

# Configuracao numerica explicita que reproduz a execucao que convergiu. Com
# optim_score=None, o MLEModel usa a aproximacao por diferencas finitas do
# L-BFGS-B. O escore de Harvey foi testado, mas produziu gradientes inteiramente
# NaN nesta especificacao com lacunas e inicializacao difusa exata.
# A convergencia nativa pode ocorrer por reducao da funcao mesmo com gradiente
# alto. Por isso esta versao tambem exige gradiente pequeno, covariancia bem
# condicionada e repeticao da mesma solucao em inicializacoes independentes.
LBFGS_PGTOL = 1e-5
LBFGS_FACTR = 1e4
LBFGS_MEMORY = 10
LBFGS_MAXFUN = 30000
LBFGS_FINITE_DIFFERENCE_EPSILON = 1e-5
GRADIENT_ACCEPTANCE_TOLERANCE = 1e-3
# O numero de condicao da covariancia bruta depende das unidades dos
# parametros (variancias, coeficientes beta e AR). Ele e mantido para auditoria,
# mas a decisao numerica usa a matriz de correlacao dos parametros, que remove
# esse efeito puramente de escala.
COVARIANCE_CORRELATION_CONDITION_MAX = 1e12
LOGLIK_CLUSTER_TOLERANCE = 0.1
PARAMETER_CLUSTER_RELATIVE_TOLERANCE = 0.02
PARAMETER_CLUSTER_ABSOLUTE_TOLERANCE = 1e-3

# Escalas das perturbacoes no espaco nao restrito. Quando existe warm start,
# uma inicializacao e exatamente mapeada, uma usa o default do statsmodels e as
# outras tres sao perturbacoes independentes da solucao mapeada.
DEFAULT_START_PERTURBATION_SCALES = [0.03, 0.08, 0.15, 0.25]
WARM_START_PERTURBATION_SCALES = [0.015, 0.04, 0.10]

# Validacao interna do Kalman por mascaramento artificial somente em 2022-2023.
# Nao e a validacao Q50.
RUN_MASKED_VALIDATION = True
VALIDATION_REQUIRED_SUCCESSES = 5
VALIDATION_MAX_ATTEMPTS = 12
VALIDATION_FRACTION = 0.03
VALIDATION_MAXITER = 1500
VALIDATION_MIN_BLOCKS_PER_DURATION_CLASS = 1
VALIDATION_DURATION_CLASSES = [
    (1, 1, "1h"),
    (2, 3, "2-3h"),
    (4, 6, "4-6h"),
    (7, 12, "7-12h"),
    (13, 24, "13-24h"),
    (25, 48, "25-48h"),
    (49, 72, "49-72h"),
]

# Mantem o objeto statsmodels completo para auditoria/reuso.
SAVE_MODEL_PICKLE = True

ALPHA_INTERVAL = 0.05  # intervalo de 95%


# =============================================================================
# 3. FUNCOES AUXILIARES
# =============================================================================

def sha256_file(path, chunk_size=1024 * 1024):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def json_safe(value):
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and not np.isfinite(value):
            return str(value)
        return value
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value) if np.isfinite(value) else str(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, pd.Series):
        return value.to_dict()
    if isinstance(value, pd.DataFrame):
        return value.to_dict(orient="records")
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    return str(value)


def transform_flow(q):
    q = pd.Series(q, copy=True, dtype=float)
    invalid_nonpositive = q.notna() & (q <= 0)
    if invalid_nonpositive.any():
        examples = q.index[invalid_nonpositive][:10].tolist()
        raise ValueError(
            "Foram encontradas vazoes observadas <= 0. O QC deve converte-las "
            f"em NaN antes do Kalman. Exemplos de indices: {examples}"
        )
    return np.log(q) if USE_LOG_FLOW else q


def engineer_precipitation_features(frame):
    """Recria as covariaveis de chuva de curto prazo do modelo final."""
    result = frame.copy()
    precipitation = pd.to_numeric(result["Precipitação"], errors="coerce")

    if precipitation.isna().any():
        raise ValueError(
            "A coluna Precipitação contem NaNs ou valores nao numericos. "
            "O Kalman nao imputara covariaveis meteorologicas."
        )
    if (precipitation < 0).any():
        raise ValueError("Foram encontradas precipitacoes negativas.")

    engineered = pd.DataFrame(index=result.index)
    engineered["precip_lag_1h"] = precipitation.shift(1, fill_value=0.0)
    engineered["precip_lag_2h"] = precipitation.shift(2, fill_value=0.0)
    engineered["tempo_sem_chuva"] = (
        precipitation.eq(0)
        .groupby(precipitation.gt(0).cumsum())
        .cumsum()
        .astype(float)
    )

    preexisting = [
        column for column in ENGINEERED_EXOG_COLUMNS if column in result.columns
    ]
    for column in preexisting:
        supplied = pd.to_numeric(result[column], errors="coerce")
        expected = engineered[column]
        if supplied.isna().any() or not np.allclose(
            supplied.to_numpy(dtype=float),
            expected.to_numpy(dtype=float),
            rtol=1e-10,
            atol=1e-10,
        ):
            raise ValueError(
                f"A coluna derivada preexistente '{column}' nao coincide com "
                "a definicao metodologica congelada. Remova-a do arquivo de "
                "entrada ou revise explicitamente a definicao."
            )

    for column in ENGINEERED_EXOG_COLUMNS:
        result[column] = engineered[column]

    return result, preexisting


def inverse_values(values):
    arr = np.asarray(values, dtype=float)
    if USE_LOG_FLOW:
        # Protecao apenas contra overflow numerico; nao e limite hidraulico.
        arr = np.exp(np.clip(arr, -50, 50))
    return arr


def apply_operational_bounds(values, lower_bound, upper_bound):
    raw = np.asarray(values, dtype=float)
    bounded = np.clip(raw, lower_bound, upper_bound)
    lower_flag = np.isfinite(raw) & (raw < lower_bound)
    upper_flag = np.isfinite(raw) & (raw > upper_bound)
    return bounded, lower_flag, upper_flag


def prediction_to_original_scale(prediction_result, lower_bound, upper_bound):
    mean_model_scale = np.asarray(prediction_result.predicted_mean, dtype=float).reshape(-1)
    ci_model_scale = np.asarray(
        prediction_result.conf_int(alpha=ALPHA_INTERVAL), dtype=float
    )
    if ci_model_scale.ndim != 2 or ci_model_scale.shape[1] < 2:
        raise RuntimeError("Formato inesperado do intervalo retornado pelo statsmodels.")

    estimate_raw = inverse_values(mean_model_scale)
    lower_raw = inverse_values(ci_model_scale[:, 0])
    upper_raw = inverse_values(ci_model_scale[:, 1])

    estimate, clipped_low, clipped_high = apply_operational_bounds(
        estimate_raw, lower_bound, upper_bound
    )
    lower = np.clip(lower_raw, lower_bound, upper_bound)
    upper = np.clip(upper_raw, lower_bound, upper_bound)

    return {
        "model_scale": mean_model_scale,
        "estimate_raw": estimate_raw,
        "lower_raw": lower_raw,
        "upper_raw": upper_raw,
        "estimate": estimate,
        "lower": lower,
        "upper": upper,
        "clipped_low": clipped_low,
        "clipped_high": clipped_high,
    }


def build_model(endog, exog, model_spec):
    ar_order = int(model_spec.get("ar_order", 0))
    return UnobservedComponents(
        endog=endog,
        exog=exog,
        level=True,
        trend=bool(model_spec["trend"]),
        stochastic_level=True,
        stochastic_trend=bool(model_spec["stochastic_trend"]),
        irregular=bool(model_spec["irregular"]),
        autoregressive=ar_order if ar_order > 0 else None,
        mle_regression=True,
        use_exact_diffuse=USE_EXACT_DIFFUSE,
    )


def observation_equation_for_spec(model_spec):
    equation = "log(Q_t) = level_t + beta' x_t"
    if int(model_spec.get("ar_order", 0)) > 0:
        equation += " + ar_t"
    return equation + " + epsilon_t" if model_spec["irregular"] else equation


def state_equations_for_spec(model_spec):
    if model_spec["trend"]:
        level_equation = "level_(t+1) = level_t + trend_t + eta_level,t"
        trend_equation = (
            "trend_(t+1) = trend_t + eta_trend,t"
            if model_spec["stochastic_trend"]
            else "trend_(t+1) = trend_t (deterministic slope)"
        )
        equations = [level_equation, trend_equation]
    else:
        equations = ["level_(t+1) = level_t + eta_level,t"]
    if int(model_spec.get("ar_order", 0)) > 0:
        equations.append("ar_(t+1) = phi * ar_t + eta_ar,t")
    return equations


def optimizer_projected_gradient_norm(result):
    retvals = getattr(result, "mle_retvals", {}) or {}
    gradient = retvals.get("gopt", None)
    if gradient is None:
        return np.nan
    gradient = np.asarray(gradient, dtype=float)
    finite_gradient = np.abs(gradient[np.isfinite(gradient)])
    return float(np.max(finite_gradient)) if finite_gradient.size else np.nan


def covariance_diagnostics(result):
    """Condicionamento bruto e livre de escala da covariancia OPG."""
    try:
        covariance = np.asarray(result.cov_params(), dtype=float)
    except Exception:
        covariance = np.array([], dtype=float)

    diagnostics = {
        "raw_condition_number": np.inf,
        "correlation_condition_number": np.inf,
        "positive_finite_diagonal": False,
        "rank": 0,
    }
    if (
        covariance.ndim != 2
        or covariance.shape[0] != covariance.shape[1]
        or covariance.size == 0
        or not np.isfinite(covariance).all()
    ):
        return diagnostics

    try:
        diagnostics["raw_condition_number"] = float(np.linalg.cond(covariance))
        diagnostics["rank"] = int(np.linalg.matrix_rank(covariance))
    except np.linalg.LinAlgError:
        return diagnostics

    diagonal = np.diag(covariance)
    positive_finite = bool(
        np.isfinite(diagonal).all() and np.all(diagonal > 0.0)
    )
    diagnostics["positive_finite_diagonal"] = positive_finite
    if not positive_finite:
        return diagnostics

    standard_errors = np.sqrt(diagonal)
    correlation = covariance / np.outer(standard_errors, standard_errors)
    correlation = 0.5 * (correlation + correlation.T)
    if not np.isfinite(correlation).all():
        return diagnostics
    try:
        diagnostics["correlation_condition_number"] = float(
            np.linalg.cond(correlation)
        )
    except np.linalg.LinAlgError:
        pass
    return diagnostics


def covariance_condition_number(result):
    """Alias historico: retorna o numero de condicao bruto para auditoria."""
    return covariance_diagnostics(result)["raw_condition_number"]


def covariance_correlation_condition_number(result):
    """Numero de condicao usado na regra de aceitacao, livre de escala."""
    return covariance_diagnostics(result)["correlation_condition_number"]


def near_zero_estimated_variances(result, threshold=1e-10):
    flagged = []
    for name, value in zip(result.model.param_names, result.params):
        if name.startswith("sigma2.") and float(value) < threshold:
            flagged.append(name)
    return flagged


def relative_parameter_distance(params, reference):
    """Distancia normalizada equivalente ao criterio atol + rtol.

    A versao anterior dividia cada diferenca apenas pela magnitude do parametro
    de referencia. Isso penalizava indevidamente coeficientes proximos de zero:
    diferencas absolutas despreziveis produziam grandes erros relativos. Aqui,
    valores <= 1 indicam equivalencia segundo tolerancias absoluta e relativa
    declaradas explicitamente.
    """
    params = np.asarray(params, dtype=float)
    reference = np.asarray(reference, dtype=float)
    scale = (
        PARAMETER_CLUSTER_ABSOLUTE_TOLERANCE
        + PARAMETER_CLUSTER_RELATIVE_TOLERANCE
        * np.maximum(np.abs(params), np.abs(reference))
    )
    return float(np.max(np.abs(params - reference) / scale))


def parameter_vector_as_json(result):
    return json.dumps(
        {
            str(name): float(value)
            for name, value in zip(result.model.param_names, result.params)
        },
        ensure_ascii=False,
        sort_keys=True,
    )


def map_transformed_parameters_by_name(source_result, target_model):
    """Transporta somente parametros homonimos para um modelo aninhado."""
    target = np.asarray(target_model.start_params, dtype=float).copy()
    source_map = {
        str(name): float(value)
        for name, value in zip(
            source_result.model.param_names,
            source_result.params,
        )
    }
    mapped_names = []
    for position, name in enumerate(target_model.param_names):
        if name in source_map and np.isfinite(source_map[name]):
            target[position] = source_map[name]
            mapped_names.append(str(name))
    if not np.isfinite(target).all():
        raise RuntimeError("Warm start mapeado contem parametro nao finito.")
    return target, mapped_names


def build_start_schedule(base_model, rng, warm_start_result=None):
    base_transformed = np.asarray(base_model.start_params, dtype=float)
    schedule = []

    if warm_start_result is not None:
        warm_transformed, mapped_names = map_transformed_parameters_by_name(
            warm_start_result,
            base_model,
        )
        schedule.append(
            {
                "params": warm_transformed,
                "transformed": True,
                "origin": "mapped_boundary_warm_start",
                "perturbation_scale": 0.0,
                "mapped_parameter_names": " | ".join(mapped_names),
            }
        )
        schedule.append(
            {
                "params": base_transformed,
                "transformed": True,
                "origin": "statsmodels_default",
                "perturbation_scale": 0.0,
                "mapped_parameter_names": "",
            }
        )
        warm_untransformed = np.asarray(
            base_model.untransform_params(warm_transformed),
            dtype=float,
        )
        for scale in WARM_START_PERTURBATION_SCALES:
            schedule.append(
                {
                    "params": warm_untransformed
                    + rng.normal(0.0, scale, size=warm_untransformed.shape),
                    "transformed": False,
                    "origin": "perturbed_mapped_warm_start",
                    "perturbation_scale": float(scale),
                    "mapped_parameter_names": " | ".join(mapped_names),
                }
            )
    else:
        schedule.append(
            {
                "params": base_transformed,
                "transformed": True,
                "origin": "statsmodels_default",
                "perturbation_scale": 0.0,
                "mapped_parameter_names": "",
            }
        )
        base_untransformed = np.asarray(
            base_model.untransform_params(base_transformed),
            dtype=float,
        )
        for scale in DEFAULT_START_PERTURBATION_SCALES:
            schedule.append(
                {
                    "params": base_untransformed
                    + rng.normal(0.0, scale, size=base_untransformed.shape),
                    "transformed": False,
                    "origin": "perturbed_statsmodels_default",
                    "perturbation_scale": float(scale),
                    "mapped_parameter_names": "",
                }
            )

    if len(schedule) != N_STARTS_PER_CANDIDATE:
        raise RuntimeError(
            "A agenda de inicializacoes nao coincide com "
            f"N_STARTS_PER_CANDIDATE={N_STARTS_PER_CANDIDATE}."
        )
    return schedule


def fit_candidate_multistart(
    endog,
    exog,
    model_spec,
    candidate_index,
    warm_start_result=None,
    warm_start_note="",
):
    base_model = build_model(endog, exog, model_spec)
    rng = np.random.default_rng(RANDOM_SEED + 100 * candidate_index)
    start_schedule = build_start_schedule(
        base_model,
        rng,
        warm_start_result=warm_start_result,
    )

    attempts = []
    fitted_records = []

    for start_id, start_config in enumerate(start_schedule):
        try:
            model = build_model(endog, exog, model_spec)

            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                warnings.filterwarnings(
                    "ignore",
                    message=r"datetime\.datetime\.utcnow\(\).*deprecated.*",
                )
                result = model.fit(
                    start_params=start_config["params"],
                    transformed=start_config["transformed"],
                    method=OPTIMIZER,
                    maxiter=MAXITER_FINAL,
                    disp=False,
                    full_output=True,
                    cov_type="opg",
                    low_memory=False,
                    optim_score=OPTIM_SCORE_METHOD,
                    pgtol=LBFGS_PGTOL,
                    factr=LBFGS_FACTR,
                    m=LBFGS_MEMORY,
                    maxfun=LBFGS_MAXFUN,
                    epsilon=LBFGS_FINITE_DIFFERENCE_EPSILON,
                )

            retvals = getattr(result, "mle_retvals", {}) or {}
            converged = bool(retvals.get("converged", False))
            warnflag = retvals.get("warnflag", np.nan)
            projected_gradient_norm = optimizer_projected_gradient_norm(result)
            finite_params = bool(np.isfinite(np.asarray(result.params)).all())
            finite_likelihood = bool(np.isfinite(float(result.llf)))
            covariance = covariance_diagnostics(result)
            covariance_raw_condition = covariance["raw_condition_number"]
            covariance_correlation_condition = covariance[
                "correlation_condition_number"
            ]
            native_ok = converged and warnflag == 0
            gradient_ok = (
                np.isfinite(projected_gradient_norm)
                and projected_gradient_norm <= GRADIENT_ACCEPTANCE_TOLERANCE
            )
            numerical_ok = (
                native_ok
                and gradient_ok
                and finite_params
                and finite_likelihood
            )
            covariance_ok = (
                covariance["positive_finite_diagonal"]
                and np.isfinite(covariance_correlation_condition)
                and covariance_correlation_condition
                <= COVARIANCE_CORRELATION_CONDITION_MAX
            )
            boundary_variances = near_zero_estimated_variances(result)
            boundary_ok = not boundary_variances
            acceptable = (
                numerical_ok
                and covariance_ok
                and boundary_ok
            )

            attempt = {
                "model_id": model_spec["model_id"],
                "ar_order": int(model_spec.get("ar_order", 0)),
                "start_id": start_id,
                "start_origin": start_config["origin"],
                "perturbation_scale": start_config["perturbation_scale"],
                "warm_start_note": warm_start_note,
                "mapped_parameter_names": start_config[
                    "mapped_parameter_names"
                ],
                "status": "ok",
                "converged": converged,
                "native_ok": native_ok,
                "gradient_ok": gradient_ok,
                "numerical_ok": numerical_ok,
                "covariance_ok": covariance_ok,
                "boundary_ok": boundary_ok,
                "boundary_variance_parameters": " | ".join(
                    boundary_variances
                ),
                "acceptable": acceptable,
                "llf": float(result.llf),
                "aic": float(result.aic),
                "bic": float(result.bic),
                "projected_gradient_inf_norm": projected_gradient_norm,
                "covariance_raw_condition_number": covariance_raw_condition,
                "covariance_correlation_condition_number": (
                    covariance_correlation_condition
                ),
                "covariance_positive_finite_diagonal": covariance[
                    "positive_finite_diagonal"
                ],
                "covariance_rank": covariance["rank"],
                "parameter_count": int(len(result.params)),
                "parameter_vector_json": parameter_vector_as_json(result),
                "iterations": retvals.get("iterations", np.nan),
                "function_calls": retvals.get("fcalls", np.nan),
                "warnflag": warnflag,
                "warnings": " | ".join(str(w.message) for w in caught),
            }
            attempts.append(attempt)
            print(
                f"  start={start_id} origin={start_config['origin']} "
                f"converged={converged} warnflag={warnflag} "
                f"grad={projected_gradient_norm:.3e} "
                f"cond_corr={covariance_correlation_condition:.3e} "
                f"boundary={boundary_variances or 'none'} "
                f"acceptable={acceptable}"
            )
            fitted_records.append(
                {
                    "result": result,
                    "attempt": attempt,
                    "numerical_ok": numerical_ok,
                    "acceptable": acceptable,
                    "boundary_variances": boundary_variances,
                }
            )

        except Exception as exc:
            attempts.append(
                {
                    "model_id": model_spec["model_id"],
                    "ar_order": int(model_spec.get("ar_order", 0)),
                    "start_id": start_id,
                    "start_origin": start_config["origin"],
                    "perturbation_scale": start_config["perturbation_scale"],
                    "warm_start_note": warm_start_note,
                    "mapped_parameter_names": start_config[
                        "mapped_parameter_names"
                    ],
                    "status": "error",
                    "converged": False,
                    "native_ok": False,
                    "gradient_ok": False,
                    "numerical_ok": False,
                    "covariance_ok": False,
                    "boundary_ok": False,
                    "boundary_variance_parameters": "",
                    "acceptable": False,
                    "llf": np.nan,
                    "aic": np.nan,
                    "bic": np.nan,
                    "projected_gradient_inf_norm": np.nan,
                    "covariance_raw_condition_number": np.inf,
                    "covariance_correlation_condition_number": np.inf,
                    "covariance_positive_finite_diagonal": False,
                    "covariance_rank": 0,
                    "parameter_count": int(len(base_model.param_names)),
                    "parameter_vector_json": "",
                    "iterations": np.nan,
                    "function_calls": np.nan,
                    "warnflag": np.nan,
                    "warnings": repr(exc),
                }
            )
            print(f"  start={start_id} ERRO: {exc!r}")

    numerical_records = [
        record for record in fitted_records if record["numerical_ok"]
    ]
    acceptable_records = [
        record for record in fitted_records if record["acceptable"]
    ]
    best_numerical_record = (
        max(numerical_records, key=lambda item: float(item["result"].llf))
        if numerical_records
        else None
    )
    numerical_cluster_count = 0
    stable_acceptable_cluster_count = 0
    robust_candidate = False
    rejection_reason = ""

    if best_numerical_record is not None:
        best_result = best_numerical_record["result"]
        best_llf = float(best_result.llf)
        best_params = np.asarray(best_result.params, dtype=float)
        cluster_records = [
            record
            for record in numerical_records
            if (
                best_llf - float(record["result"].llf)
                <= LOGLIK_CLUSTER_TOLERANCE
                and relative_parameter_distance(
                    record["result"].params,
                    best_params,
                )
                <= 1.0
            )
        ]
        numerical_cluster_count = len(cluster_records)
        stable_acceptable_cluster_count = sum(
            bool(record["acceptable"]) for record in cluster_records
        )
        best_is_acceptable = bool(best_numerical_record["acceptable"])
        robust_candidate = (
            best_is_acceptable
            and stable_acceptable_cluster_count >= MIN_STABLE_STARTS
        )
        if not best_is_acceptable:
            failed_parts = []
            attempt = best_numerical_record["attempt"]
            if not bool(attempt["covariance_ok"]):
                failed_parts.append("covariancia/correlacao mal condicionada")
            if not bool(attempt["boundary_ok"]):
                failed_parts.append(
                    "variancia no limite: "
                    + str(attempt["boundary_variance_parameters"])
                )
            rejection_reason = (
                "A melhor bacia de verossimilhanca foi numericamente "
                "convergente, mas nao e uma solucao interior aceitavel ("
                + "; ".join(failed_parts or ["criterio nao satisfeito"])
                + "). O resultado deve ser transferido ao modelo aninhado."
            )
        elif stable_acceptable_cluster_count < MIN_STABLE_STARTS:
            rejection_reason = (
                f"Somente {stable_acceptable_cluster_count} inicializacoes "
                "aceitaveis reproduziram a melhor solucao; minimo exigido="
                f"{MIN_STABLE_STARTS}."
            )
    else:
        rejection_reason = (
            "Nenhuma inicializacao satisfez conjuntamente convergencia "
            "nativa, gradiente e finitude."
        )

    best_result = (
        best_numerical_record["result"]
        if best_numerical_record is not None
        else None
    )
    best_covariance = (
        covariance_diagnostics(best_result)
        if best_result is not None
        else {
            "raw_condition_number": np.inf,
            "correlation_condition_number": np.inf,
        }
    )
    summary = {
        "model_id": model_spec["model_id"],
        "trend": bool(model_spec["trend"]),
        "stochastic_trend": bool(model_spec["stochastic_trend"]),
        "ar_order": int(model_spec.get("ar_order", 0)),
        "irregular_R_estimated": bool(model_spec["irregular"]),
        "warm_start_source": model_spec.get("warm_start_source"),
        "warm_start_note": warm_start_note,
        "starts_total": N_STARTS_PER_CANDIDATE,
        "starts_numerically_valid": int(len(numerical_records)),
        "starts_acceptable": int(len(acceptable_records)),
        "numerical_cluster_count": int(numerical_cluster_count),
        "stable_acceptable_cluster_count": int(
            stable_acceptable_cluster_count
        ),
        "robust_candidate": bool(robust_candidate),
        "best_llf": float(best_result.llf) if best_result is not None else np.nan,
        "best_aic": float(best_result.aic) if best_result is not None else np.nan,
        "best_bic": float(best_result.bic) if best_result is not None else np.nan,
        "best_gradient_inf_norm": (
            optimizer_projected_gradient_norm(best_result)
            if best_result is not None
            else np.nan
        ),
        "best_covariance_condition_number": (
            best_covariance["raw_condition_number"]
        ),
        "best_covariance_correlation_condition_number": (
            best_covariance["correlation_condition_number"]
        ),
        "best_boundary_variance_parameters": (
            " | ".join(best_numerical_record["boundary_variances"])
            if best_numerical_record is not None
            else ""
        ),
        "rejection_reason": rejection_reason,
    }
    print(pd.DataFrame([summary]).to_string(index=False))
    return (
        best_result if robust_candidate else None,
        pd.DataFrame(attempts),
        summary,
        fitted_records,
    )


def choose_warm_start_result(source_records, required_boundary_parameters):
    numerical_records = [
        record for record in source_records if record["numerical_ok"]
    ]
    if not numerical_records:
        return None, "fonte sem solucao numericamente valida"

    required = set(required_boundary_parameters or [])
    matching = [
        record
        for record in numerical_records
        if required.issubset(set(record["boundary_variances"]))
    ]
    pool = matching if matching else numerical_records
    chosen = max(pool, key=lambda item: float(item["result"].llf))
    source_attempt = chosen["attempt"]
    match_text = "fronteira correspondente" if matching else "fallback numerico"
    note = (
        f"{match_text}; model={source_attempt['model_id']}; "
        f"start={source_attempt['start_id']}; llf={source_attempt['llf']:.12f}; "
        "boundary="
        f"{source_attempt['boundary_variance_parameters'] or 'none'}"
    )
    return chosen["result"], note


def fit_and_select_model(endog, exog):
    candidate_results = []
    attempt_tables = []
    candidate_summaries = []
    records_by_model = {}

    for candidate_index, model_spec in enumerate(CANDIDATE_MODELS):
        print("\nAjustando candidato:", model_spec["model_id"])
        warm_start_result = None
        warm_start_note = "sem warm start"
        source_model_id = model_spec.get("warm_start_source")
        if source_model_id is not None:
            warm_start_result, warm_start_note = choose_warm_start_result(
                records_by_model.get(source_model_id, []),
                model_spec.get("warm_start_boundary_parameters", []),
            )
            if warm_start_result is None:
                raise RuntimeError(
                    f"Nao foi possivel construir warm start para "
                    f"{model_spec['model_id']} a partir de {source_model_id}."
                )
            print("  warm start:", warm_start_note)

        result, attempts_df, summary, fitted_records = fit_candidate_multistart(
            endog,
            exog,
            model_spec,
            candidate_index,
            warm_start_result=warm_start_result,
            warm_start_note=warm_start_note,
        )
        records_by_model[model_spec["model_id"]] = fitted_records
        attempt_tables.append(attempts_df)
        candidate_summaries.append(summary)
        if result is not None:
            candidate_results.append((model_spec, result))

    attempts_all_df = pd.concat(attempt_tables, ignore_index=True)
    selection_df = pd.DataFrame(candidate_summaries)
    if not candidate_results:
        selection_failure_path = OUTPUT_DIR / "KALMAN_MODEL_SELECTION_FAILURES.csv"
        attempts_failure_path = OUTPUT_DIR / "KALMAN_OPTIMIZATION_FAILURES.csv"
        selection_df.to_csv(selection_failure_path, index=False)
        attempts_all_df.to_csv(attempts_failure_path, index=False)
        print(selection_df.to_string(index=False))
        print(attempts_all_df.to_string(index=False))
        raise RuntimeError(
            "Nenhum modelo candidato apresentou convergencia robusta em pelo "
            "menos tres inicializacoes. Diagnosticos salvos em "
            f"{selection_failure_path} e {attempts_failure_path}. "
            "Nao use esta execucao no manuscrito."
        )

    selected_spec, selected_result = min(
        candidate_results,
        key=lambda item: float(item[1].bic),
    )
    selection_df["selected_by_bic"] = (
        selection_df["model_id"] == selected_spec["model_id"]
    )
    return (
        selected_spec,
        selected_result.model,
        selected_result,
        attempts_all_df,
        selection_df,
    )


def consecutive_gap_table(mask, index, method_by_row):
    mask = np.asarray(mask, dtype=bool)
    rows = []
    gap_id_by_row = np.full(len(mask), np.nan)
    gap_id = 0
    i = 0
    while i < len(mask):
        if not mask[i]:
            i += 1
            continue
        start = i
        current_method = method_by_row[i]
        while (
            i + 1 < len(mask)
            and mask[i + 1]
            and method_by_row[i + 1] == current_method
        ):
            i += 1
        end = i
        gap_id += 1
        gap_id_by_row[start : end + 1] = gap_id

        if start == 0:
            position = "borda_inicial_serie"
        elif end == len(mask) - 1:
            position = "borda_final_serie"
        elif index[start].year != index[end].year:
            position = "cruza_ano"
        else:
            position = "interior"

        rows.append(
            {
                "gap_id": gap_id,
                "inicio": index[start],
                "fim": index[end],
                "duracao_horas": end - start + 1,
                "posicao": position,
                "metodo": current_method,
                "ano_inicio": int(index[start].year),
                "ano_fim": int(index[end].year),
            }
        )
        i += 1
    return pd.DataFrame(rows), gap_id_by_row


def observed_gap_lengths(mask):
    mask = np.asarray(mask, dtype=bool)
    lengths = []
    i = 0
    while i < len(mask):
        if not mask[i]:
            i += 1
            continue
        start = i
        while i + 1 < len(mask) and mask[i + 1]:
            i += 1
        lengths.append(i - start + 1)
        i += 1
    return lengths


def duration_class(length):
    length = int(length)
    for lower, upper, label in VALIDATION_DURATION_CLASSES:
        if lower <= length <= upper:
            return label
    return f">{VALIDATION_DURATION_CLASSES[-1][1]}h"


def generate_artificial_mask(y, real_gap_lengths, fraction, seed):
    """Cria blocos artificiais continuos com duracoes das lacunas reais."""
    rng = np.random.default_rng(seed)
    observed = y.notna().to_numpy()
    selected = np.zeros(len(y), dtype=bool)
    target = max(1, int(observed.sum() * fraction))
    lengths = [int(v) for v in real_gap_lengths if int(v) >= 1]
    if not lengths:
        lengths = [1, 2, 3, 6, 12, 24]

    blocks = []

    def try_place(length, attempts=20000):
        length = min(int(length), len(y))
        for _ in range(attempts):
            start = int(rng.integers(0, len(y) - length + 1))
            end = start + length
            # Uma hora de separacao impede que dois blocos sejam fundidos.
            guard_start = max(0, start - 1)
            guard_end = min(len(y), end + 1)
            if (
                observed[start:end].all()
                and not selected[guard_start:guard_end].any()
            ):
                selected[start:end] = True
                blocks.append(
                    {
                        "bloco_id": len(blocks) + 1,
                        "inicio_posicao": start,
                        "fim_posicao": end - 1,
                        "inicio": y.index[start],
                        "fim": y.index[end - 1],
                        "duracao_horas": length,
                        "classe_duracao": duration_class(length),
                    }
                )
                return True
        return False

    # Garante cobertura das classes de duracao realmente presentes.
    for lower, upper, label in VALIDATION_DURATION_CLASSES:
        class_lengths = [v for v in lengths if lower <= v <= upper]
        for _ in range(VALIDATION_MIN_BLOCKS_PER_DURATION_CLASS):
            if class_lengths:
                try_place(int(rng.choice(class_lengths)))

    attempts = 0
    max_attempts = 200000
    while selected.sum() < target and attempts < max_attempts:
        attempts += 1
        try_place(int(rng.choice(lengths)), attempts=1)

    if selected.sum() < max(1, int(0.8 * target)):
        raise RuntimeError(
            "Nao foi possivel construir mascara artificial suficiente para validacao."
        )
    return selected, pd.DataFrame(blocks)


def metrics(y_true, y_pred, lower=None, upper=None):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    finite = np.isfinite(y_true) & np.isfinite(y_pred)
    y_true = y_true[finite]
    y_pred = y_pred[finite]
    error = y_pred - y_true
    denominator = np.sum((y_true - np.mean(y_true)) ** 2)
    nse = 1.0 - np.sum(error**2) / denominator if denominator > 0 else np.nan
    output = {
        "n": int(len(y_true)),
        "MAE_L_s": float(np.mean(np.abs(error))),
        "RMSE_L_s": float(np.sqrt(np.mean(error**2))),
        "bias_L_s": float(np.mean(error)),
        "NSE": float(nse),
    }
    if lower is not None and upper is not None:
        lower = np.asarray(lower, dtype=float)[finite]
        upper = np.asarray(upper, dtype=float)[finite]
        output["coverage_95"] = float(
            np.mean((y_true >= lower) & (y_true <= upper))
        )
        output["mean_interval_width_L_s"] = float(np.mean(upper - lower))
    return output


def masked_validation(
    y_train_original,
    y_train_model,
    x_train,
    final_params,
    model_spec,
    lower_bound,
    upper_bound,
    reference_gap_lengths,
):
    records = []
    duration_records = []
    block_metric_records = []
    point_records = []
    failure_records = []
    successful_repeats = 0

    for attempt in range(1, VALIDATION_MAX_ATTEMPTS + 1):
        if successful_repeats >= VALIDATION_REQUIRED_SUCCESSES:
            break
        seed = RANDOM_SEED + 1000 + attempt
        artificial_mask, blocks_df = generate_artificial_mask(
            y_train_original,
            real_gap_lengths=reference_gap_lengths,
            fraction=VALIDATION_FRACTION,
            seed=seed,
        )
        y_masked = y_train_model.copy()
        y_masked.iloc[artificial_mask] = np.nan

        caught = []
        try:
            model = build_model(y_masked, x_train, model_spec)
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                warnings.filterwarnings(
                    "ignore",
                    message=r"datetime\.datetime\.utcnow\(\).*deprecated.*",
                )
                result = model.fit(
                    start_params=np.asarray(final_params, dtype=float),
                    transformed=True,
                    method=OPTIMIZER,
                    maxiter=VALIDATION_MAXITER,
                    disp=False,
                    full_output=True,
                    cov_type="opg",
                    low_memory=False,
                    optim_score=OPTIM_SCORE_METHOD,
                    pgtol=LBFGS_PGTOL,
                    factr=LBFGS_FACTR,
                    m=LBFGS_MEMORY,
                    maxfun=LBFGS_MAXFUN,
                    epsilon=LBFGS_FINITE_DIFFERENCE_EPSILON,
                )
        except Exception as exc:
            failure_records.append(
                {
                    "tentativa": attempt,
                    "semente": seed,
                    "status": "erro",
                    "converged": False,
                    "warnflag": np.nan,
                    "projected_gradient_inf_norm": np.nan,
                    "covariance_raw_condition_number": np.inf,
                    "covariance_correlation_condition_number": np.inf,
                    "boundary_variance_parameters": "",
                    "artificialmente_mascarados": int(artificial_mask.sum()),
                    "avisos": " | ".join(str(w.message) for w in caught),
                    "erro": repr(exc),
                }
            )
            print(f"Validacao tentativa={attempt}: ERRO {exc!r}")
            continue

        retvals = getattr(result, "mle_retvals", {}) or {}
        converged = bool(retvals.get("converged", False))
        warnflag = retvals.get("warnflag", np.nan)
        gradient_norm = optimizer_projected_gradient_norm(result)
        covariance = covariance_diagnostics(result)
        covariance_raw_condition = covariance["raw_condition_number"]
        covariance_correlation_condition = covariance[
            "correlation_condition_number"
        ]
        boundary_variances = near_zero_estimated_variances(result)
        acceptable = (
            converged
            and warnflag == 0
            and np.isfinite(gradient_norm)
            and gradient_norm <= GRADIENT_ACCEPTANCE_TOLERANCE
            and covariance["positive_finite_diagonal"]
            and np.isfinite(covariance_correlation_condition)
            and covariance_correlation_condition
            <= COVARIANCE_CORRELATION_CONDITION_MAX
            and not boundary_variances
            and np.isfinite(np.asarray(result.params, dtype=float)).all()
            and np.isfinite(float(result.llf))
        )
        if not acceptable:
            failure_records.append(
                {
                    "tentativa": attempt,
                    "semente": seed,
                    "status": "rejeitada_por_criterio_numerico",
                    "converged": converged,
                    "warnflag": warnflag,
                    "projected_gradient_inf_norm": gradient_norm,
                    "covariance_raw_condition_number": covariance_raw_condition,
                    "covariance_correlation_condition_number": (
                        covariance_correlation_condition
                    ),
                    "boundary_variance_parameters": " | ".join(
                        boundary_variances
                    ),
                    "artificialmente_mascarados": int(artificial_mask.sum()),
                    "avisos": " | ".join(str(w.message) for w in caught),
                    "erro": "",
                }
            )
            print(
                f"Validacao tentativa={attempt}: REJEITADA "
                f"grad={gradient_norm:.3e} "
                f"cond_corr={covariance_correlation_condition:.3e}"
            )
            continue

        successful_repeats += 1
        repeat = successful_repeats
        print(
            f"Validacao tentativa={attempt}: ACEITA como repeticao={repeat}; "
            f"grad={gradient_norm:.3e} "
            f"cond_corr={covariance_correlation_condition:.3e}"
        )

        predictions = {
            "smoother_bidirecional": result.get_prediction(
                start=0,
                end=len(y_masked) - 1,
                information_set="smoothed",
                signal_only=False,
            ),
            "filtro_causal_1_passo": result.get_prediction(
                start=0,
                end=len(y_masked) - 1,
                information_set="predicted",
                signal_only=False,
            ),
        }

        for method_name, prediction in predictions.items():
            converted = prediction_to_original_scale(
                prediction, lower_bound, upper_bound
            )
            variants = {
                "sem_limites": (
                    converted["estimate_raw"],
                    converted["lower_raw"],
                    converted["upper_raw"],
                ),
                "com_limites": (
                    converted["estimate"],
                    converted["lower"],
                    converted["upper"],
                ),
            }
            for bounds_label, (estimate, lower, upper) in variants.items():
                metric_values = metrics(
                    y_train_original.iloc[artificial_mask].to_numpy(),
                    estimate[artificial_mask],
                    lower[artificial_mask],
                    upper[artificial_mask],
                )
                records.append(
                    {
                        "repeticao": repeat,
                        "tentativa": attempt,
                        "semente": seed,
                        "metodo": method_name,
                        "aplicacao_limites": bounds_label,
                        "converged": converged,
                        "projected_gradient_inf_norm": gradient_norm,
                        "covariance_raw_condition_number": (
                            covariance_raw_condition
                        ),
                        "covariance_correlation_condition_number": (
                            covariance_correlation_condition
                        ),
                        "artificialmente_mascarados": int(artificial_mask.sum()),
                        "avisos": " | ".join(str(w.message) for w in caught),
                        **metric_values,
                    }
                )

                for class_label in blocks_df["classe_duracao"].unique():
                    class_mask = np.zeros(len(y_train_original), dtype=bool)
                    class_blocks = blocks_df.loc[
                        blocks_df["classe_duracao"] == class_label
                    ]
                    for block in class_blocks.itertuples(index=False):
                        class_mask[
                            int(block.inicio_posicao) : int(block.fim_posicao) + 1
                        ] = True
                    class_metrics = metrics(
                        y_train_original.iloc[class_mask].to_numpy(),
                        estimate[class_mask],
                        lower[class_mask],
                        upper[class_mask],
                    )
                    duration_records.append(
                        {
                            "repeticao": repeat,
                            "tentativa": attempt,
                            "semente": seed,
                            "metodo": method_name,
                            "aplicacao_limites": bounds_label,
                            "classe_duracao": class_label,
                            "numero_blocos": int(len(class_blocks)),
                            **class_metrics,
                        }
                    )

                for block in blocks_df.itertuples(index=False):
                    block_slice = slice(
                        int(block.inicio_posicao), int(block.fim_posicao) + 1
                    )
                    block_metrics = metrics(
                        y_train_original.iloc[block_slice].to_numpy(),
                        estimate[block_slice],
                        lower[block_slice],
                        upper[block_slice],
                    )
                    block_metric_records.append(
                        {
                            "repeticao": repeat,
                            "tentativa": attempt,
                            "semente": seed,
                            "metodo": method_name,
                            "aplicacao_limites": bounds_label,
                            "bloco_id": int(block.bloco_id),
                            "inicio": block.inicio,
                            "fim": block.fim,
                            "duracao_horas": int(block.duracao_horas),
                            "classe_duracao": block.classe_duracao,
                            **block_metrics,
                        }
                    )

                    for position in range(
                        int(block.inicio_posicao), int(block.fim_posicao) + 1
                    ):
                        point_records.append(
                            {
                                "repeticao": repeat,
                                "tentativa": attempt,
                                "semente": seed,
                                "metodo": method_name,
                                "aplicacao_limites": bounds_label,
                                "bloco_id": int(block.bloco_id),
                                "datetime": y_train_original.index[position],
                                "duracao_bloco_horas": int(block.duracao_horas),
                                "classe_duracao": block.classe_duracao,
                                "observado_L_s": float(
                                    y_train_original.iloc[position]
                                ),
                                "estimado_L_s": float(estimate[position]),
                                "IC95_inferior_L_s": float(lower[position]),
                                "IC95_superior_L_s": float(upper[position]),
                                "erro_estimado_menos_observado_L_s": float(
                                    estimate[position]
                                    - y_train_original.iloc[position]
                                ),
                            }
                        )

    validation_df = pd.DataFrame(records)
    failure_df = pd.DataFrame(failure_records)
    if successful_repeats < VALIDATION_REQUIRED_SUCCESSES:
        failure_path = OUTPUT_DIR / "KALMAN_VALIDATION_FAILURES.csv"
        failure_df.to_csv(failure_path, index=False)
        raise RuntimeError(
            f"A validacao obteve somente {successful_repeats} repeticoes "
            f"aceitaveis; minimo exigido={VALIDATION_REQUIRED_SUCCESSES}. "
            f"Diagnosticos salvos em {failure_path}."
        )
    return (
        validation_df,
        pd.DataFrame(duration_records),
        pd.DataFrame(block_metric_records),
        pd.DataFrame(point_records),
        failure_df,
    )


def matrix_to_long(name, matrix):
    arr = np.asarray(matrix)
    rows = []
    if arr.ndim == 0:
        rows.append({"matriz": name, "tempo": 0, "linha": 0, "coluna": 0, "valor": float(arr)})
    elif arr.ndim == 1:
        for i, value in enumerate(arr):
            rows.append({"matriz": name, "tempo": 0, "linha": i, "coluna": 0, "valor": float(value)})
    elif arr.ndim == 2:
        for i in range(arr.shape[0]):
            for j in range(arr.shape[1]):
                rows.append({"matriz": name, "tempo": 0, "linha": i, "coluna": j, "valor": float(arr[i, j])})
    elif arr.ndim == 3:
        time_slices = [0] if arr.shape[2] == 1 else [0, arr.shape[2] - 1]
        for t in sorted(set(time_slices)):
            for i in range(arr.shape[0]):
                for j in range(arr.shape[1]):
                    rows.append({"matriz": name, "tempo": t, "linha": i, "coluna": j, "valor": float(arr[i, j, t])})
    else:
        rows.append({"matriz": name, "tempo": np.nan, "linha": np.nan, "coluna": np.nan, "valor": str(arr.shape)})
    return rows


def format_excel(path):
    workbook = openpyxl.load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        for cell in worksheet[1]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="1F4E78")
            cell.alignment = openpyxl.styles.Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        worksheet.row_dimensions[1].height = 36
        for column_cells in worksheet.iter_cols(
            min_row=1,
            max_row=min(worksheet.max_row, 300),
            min_col=1,
            max_col=worksheet.max_column,
        ):
            letter = column_cells[0].column_letter
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 11), 32)
    workbook.save(path)


def standardized_innovation_diagnostics(result):
    """Resume os erros de previsao padronizados finitos do filtro."""
    values = np.asarray(
        result.filter_results.standardized_forecasts_error,
        dtype=float,
    ).reshape(result.model.k_endog, -1)[0].copy()
    burn = max(int(getattr(result, "loglikelihood_burn", 0)), result.model.k_states)
    values[:burn] = np.nan
    finite = values[np.isfinite(values)]
    if finite.size < 10:
        raise RuntimeError("Erros de previsao finitos insuficientes para diagnostico.")

    rows = [
        {"diagnostico": "n", "lag": np.nan, "estatistica": float(finite.size), "p_valor": np.nan},
        {"diagnostico": "media", "lag": np.nan, "estatistica": float(np.mean(finite)), "p_valor": np.nan},
        {"diagnostico": "desvio_padrao", "lag": np.nan, "estatistica": float(np.std(finite, ddof=1)), "p_valor": np.nan},
        {"diagnostico": "assimetria", "lag": np.nan, "estatistica": float(stats.skew(finite, bias=False)), "p_valor": np.nan},
        {"diagnostico": "curtose_excesso", "lag": np.nan, "estatistica": float(stats.kurtosis(finite, fisher=True, bias=False)), "p_valor": np.nan},
    ]
    jb = stats.jarque_bera(finite)
    rows.append(
        {
            "diagnostico": "Jarque-Bera",
            "lag": np.nan,
            "estatistica": float(jb.statistic),
            "p_valor": float(jb.pvalue),
        }
    )
    lags = [lag for lag in [1, 24, 168] if lag < finite.size]
    if lags:
        lb = acorr_ljungbox(finite, lags=lags, return_df=True)
        for lag, row in lb.iterrows():
            rows.append(
                {
                    "diagnostico": "Ljung-Box",
                    "lag": int(lag),
                    "estatistica": float(row["lb_stat"]),
                    "p_valor": float(row["lb_pvalue"]),
                }
            )
    return pd.DataFrame(rows), values


def numerical_stability_table(result):
    """Registra condicionamento e parametros proximos ao limite zero."""
    covariance = np.asarray(result.cov_params(), dtype=float)
    singular_values = np.linalg.svd(covariance, compute_uv=False)
    covariance = covariance_diagnostics(result)
    parameter_names = list(result.model.param_names)
    parameter_values = np.asarray(result.params, dtype=float)
    rows = [
        {
            "item": "covariance_raw_condition_number",
            "parametro": "",
            "valor": covariance["raw_condition_number"],
            "limite_diagnostico": np.nan,
            "flag": False,
        },
        {
            "item": "covariance_correlation_condition_number",
            "parametro": "",
            "valor": covariance["correlation_condition_number"],
            "limite_diagnostico": COVARIANCE_CORRELATION_CONDITION_MAX,
            "flag": bool(
                covariance["correlation_condition_number"]
                > COVARIANCE_CORRELATION_CONDITION_MAX
            ),
        },
        {
            "item": "covariance_positive_finite_diagonal",
            "parametro": "",
            "valor": int(covariance["positive_finite_diagonal"]),
            "limite_diagnostico": 1,
            "flag": not covariance["positive_finite_diagonal"],
        },
    ]
    for index, singular_value in enumerate(singular_values):
        rows.append(
            {
                "item": "covariance_singular_value",
                "parametro": f"s_{index + 1}",
                "valor": float(singular_value),
                "limite_diagnostico": np.nan,
                "flag": False,
            }
        )
    for name, value in zip(parameter_names, parameter_values):
        if name.startswith("sigma2."):
            rows.append(
                {
                    "item": "variance_near_zero",
                    "parametro": name,
                    "valor": float(value),
                    "limite_diagnostico": 1e-10,
                    "flag": bool(value < 1e-10),
                }
            )
    return pd.DataFrame(rows)


def write_supplementary_report(
    path,
    metadata,
    parameters_df,
    validation_summary_df,
    validation_duration_summary_df,
    gap_table_df,
):
    """Gera um texto-base auditavel para a secao suplementar do Kalman."""
    parameter_lines = "\n".join(
        f"- {row.parametro}: {row.estimativa:.10g}"
        for row in parameters_df.itertuples(index=False)
    )
    validation_lines = "Nenhuma validacao executada."
    if not validation_summary_df.empty:
        validation_lines = validation_summary_df.to_string(index=False)
    duration_lines = "Nenhuma estratificacao disponivel."
    if not validation_duration_summary_df.empty:
        duration_lines = validation_duration_summary_df.to_string(index=False)

    text = f"""KALMAN RECONSTRUCTION — SUPPLEMENTARY METHODS AND AUDIT

1. DATA PROVENANCE AND FROZEN SAMPLE
Input file: {metadata['input_file']}
Input SHA-256: {metadata['input_sha256']}
Period: {metadata['period_start']} to {metadata['period_end']} (hourly grid)
Rows: {metadata['number_rows']}
Original missing flow values: {metadata['missing_total']} ({metadata['missing_2022_2023']} in 2022–2023; {metadata['missing_2024']} in 2024)
Observed values altered by this script: 0

2. RESPONSE, COVARIATES, AND TEMPORAL AVAILABILITY
Response: hourly influent flow (L/s), modeled on the natural-log scale.
Active exogenous variables: {', '.join(metadata['exog_active'])}
All exogenous variables were standardized with means and standard deviations estimated exclusively from 2022–2023. Current-hour precipitation and meteorological measurements, short precipitation lags (1 and 2 h), dry-spell duration, and deterministic calendar descriptors were used. The previously tested 20-day precipitation accumulation and its saturation interaction were deliberately excluded from the final Kalman reconstruction because this long-memory representation could propagate antecedent wetness and inflate reconstructed flows.

3. STATE-SPACE MODEL
Selected candidate: {metadata['selected_model_id']}.
Observation equation: {metadata['observation_equation']}
State equations:
{chr(10).join('- ' + equation for equation in metadata['state_equations'])}
Implementation: {metadata['model_class']} in statsmodels {metadata['software_versions']['statsmodels']}.
State dimension: {metadata['state_dimension']}; state names: {', '.join(metadata['state_names'])}.
Initialization: {metadata['initialization']}.
Observation covariance R: {metadata['observation_covariance_R_status']}.

4. PARAMETER ESTIMATION AND NUMERICAL OPTIMIZATION
Parameters were estimated by maximum likelihood, with likelihood evaluation by the Kalman filter, using only 2022–2023. Four parsimonious nested state structures were compared after preliminary variance components reached the zero boundary. The full local-level + AR(1) + irregular model was fitted first. Boundary solutions were used only to construct parameter-name-matched warm starts for the corresponding reduced models; a boundary solution was never accepted as the final model. Each candidate used L-BFGS-B with {metadata['multi_start_attempts_per_candidate']} deterministic/reproducibly perturbed starts (seed {metadata['random_seed']}). A candidate was retained only when its highest-likelihood numerically valid basin was interior, had a finite positive covariance diagonal, satisfied the scale-free parameter-correlation conditioning criterion, and was reproduced by at least {metadata['minimum_stable_starts']} starts; the lowest-BIC robust candidate was selected. The final projected-gradient infinity norm was {metadata['projected_gradient_inf_norm']}. The raw OPG covariance condition number was {metadata['covariance_raw_condition_number']} (reported for audit only), whereas the scale-free parameter-correlation condition number used for acceptance was {metadata['covariance_correlation_condition_number']}. Log-likelihood={metadata['log_likelihood']}; AIC={metadata['aic']}; BIC={metadata['bic']}.

Estimated parameters:
{parameter_lines}

The exact design, transition, selection, process-covariance (Q or equivalent), observation-covariance (R or equivalent), initial-state, and initial-covariance matrices are stored in the 'matrizes_estado' worksheet of KALMAN_RECONSTRUCAO_FINAL.xlsx. Q, R, and the initial covariance are expressed on the log-flow model scale. Candidate comparison is stored in 'selecao_modelo'; all starts are stored in 'otimizacao'; complete singular values and covariance-conditioning diagnostics are stored in 'estabilidade_numerica'.

5. RECONSTRUCTION RULE
Missing values in 2022–2023 were reconstructed with the fixed-interval Kalman smoother, which conditions on observations on both sides of an internal gap within the development period. Parameters were then frozen. Missing values in 2024 were reconstructed sequentially with one-step-ahead Kalman predictions: each estimate used response observations only through the preceding hour and exogenous information available for the hour being reconstructed. Later 2024 observations were never used to revise an earlier 2024 imputation. Original non-missing flow observations were preserved exactly.

The point estimate on the original scale is exp(mu), i.e., the conditional median under the Gaussian log-scale model; the 95% interval endpoints are the exponentiated Gaussian prediction limits. Only reconstructed values were constrained to the technically defined range {metadata['lower_operational_bound_L_s']}–{metadata['upper_operational_bound_L_s']} L/s. The lower threshold is a technical-operational plausibility bound supported by historical dry-weather behavior; the upper threshold equals the largest originally valid observed flow. Counts affected by these bounds were {metadata['clipped_lower_missing_count']} (lower) and {metadata['clipped_upper_missing_count']} (upper), corresponding to {metadata['clipped_2022_2023_count']} reconstructed values in 2022–2023 and {metadata['clipped_2024_count']} in 2024. Validation metrics are reported both before and after applying these bounds.

6. GAP STRUCTURE
Number of gaps: {len(gap_table_df)}
Maximum full-series gap: {metadata['max_gap_hours']} h
Maximum 2022–2023 gap: {metadata['max_gap_hours_2022_2023']} h
Maximum 2024 gap: {metadata['max_gap_hours_2024']} h
The complete gap inventory, including start, end, duration, location, year, and reconstruction method, is stored in the 'lacunas' worksheet.
Series-edge gaps present: {metadata['series_edge_gaps_present']}.

7. ARTIFICIAL-GAP VALIDATION
Validation was performed only on originally observed 2022–2023 responses and is separate from the subsequent Q50 consistency assessment. Artificial blocks were placed only in 2022–2023, but their lengths were sampled from the complete 2022–2024 empirical gap-length distribution so that every duration class affecting the final reconstruction, including 25–48 h, was assessed without using any missing 2024 response value. The run required {metadata['masked_validation_required_successes']} numerically acceptable refits and recorded {metadata['masked_validation_successes']} successes and {metadata['masked_validation_attempt_failures']} rejected attempts. Both fixed-interval smoother and causal one-step filter reconstructions were evaluated with and without operational limits. Bias is defined as estimate minus observation.

Overall results:
{validation_lines}

Results by gap-duration class:
{duration_lines}

Timestamp-level artificial-gap predictions and block-level metrics are retained in the workbook for audit and figure preparation.

8. SOFTWARE, REPRODUCIBILITY, AND FILE MAP
Python: {metadata['software_versions']['python']}
NumPy: {metadata['software_versions']['numpy']}; pandas: {metadata['software_versions']['pandas']}; SciPy: {metadata['software_versions']['scipy']}; statsmodels: {metadata['software_versions']['statsmodels']}; openpyxl: {metadata['software_versions']['openpyxl']}.
The JSON metadata file records all settings and software versions. The model pickle preserves the fitted 2022–2023 statsmodels result. When executed from the supplied .py file, an exact copy of the executed code is included in the package. The run manifest records SHA-256 hashes for every deliverable. Q50 validation is intentionally outside this Kalman-only script and must remain a distinct subsequent consistency-assessment step in the manuscript and Supplementary Material.
"""
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(text)


# =============================================================================
# 4. LEITURA E VALIDACOES ESTRUTURAIS
# =============================================================================

if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"Arquivo nao encontrado: {INPUT_FILE}\n"
        "Confirme que ele esta diretamente em MyDrive/Colab Notebooks/."
    )

input_sha256 = sha256_file(INPUT_FILE)
df = pd.read_excel(INPUT_FILE, sheet_name=INPUT_SHEET, engine="openpyxl")

# Nesta etapa sao exigidas somente as colunas que precisam existir na base
# pre-Kalman. As tres covariaveis derivadas serao criadas abaixo.
required_input_columns = [DATETIME_COLUMN, FLOW_COLUMN, *RAW_EXOG_COLUMNS]
missing_columns = [
    column for column in required_input_columns if column not in df.columns
]
if missing_columns:
    raise KeyError(
        "Colunas de entrada obrigatorias ausentes:\n- "
        + "\n- ".join(missing_columns)
        + "\n\nColunas encontradas:\n- "
        + "\n- ".join(map(str, df.columns.tolist()))
    )

df[DATETIME_COLUMN] = pd.to_datetime(df[DATETIME_COLUMN], errors="raise")
df = df.sort_values(DATETIME_COLUMN, kind="stable").reset_index(drop=True)

if df[DATETIME_COLUMN].duplicated().any():
    duplicated = df.loc[df[DATETIME_COLUMN].duplicated(False), DATETIME_COLUMN]
    raise ValueError(
        "Existem timestamps duplicados. Exemplos:\n"
        + duplicated.head(20).to_string(index=False)
    )

expected_index = pd.date_range(
    df[DATETIME_COLUMN].min(), df[DATETIME_COLUMN].max(), freq="h"
)
actual_index = pd.DatetimeIndex(df[DATETIME_COLUMN])
if not actual_index.equals(expected_index):
    missing_timestamps = expected_index.difference(actual_index)
    raise ValueError(
        "A grade temporal nao e horaria e continua. O codigo nao criara linhas "
        "silenciosamente. Timestamps ausentes (primeiros 20):\n"
        + "\n".join(map(str, missing_timestamps[:20]))
    )

# A engenharia so ocorre depois da ordenacao e da confirmacao de uma grade
# horaria continua, evitando lags ou janelas deslocados por linhas ausentes.
df, preexisting_engineered_columns = engineer_precipitation_features(df)

missing_after_engineering = [
    column for column in EXOG_COLUMNS if column not in df.columns
]
if missing_after_engineering:
    raise RuntimeError(
        "Falha interna na engenharia das covariaveis:\n- "
        + "\n- ".join(missing_after_engineering)
    )

outside_period = ~df[DATETIME_COLUMN].between(TRAIN_START, TEST_END)
if outside_period.any():
    raise ValueError(
        "Foram encontradas linhas fora do periodo congelado 2022-2024. "
        "Revise TRAIN_START/TEST_END antes de prosseguir."
    )

train_mask = df[DATETIME_COLUMN].between(TRAIN_START, TRAIN_END).to_numpy()
test_mask = df[DATETIME_COLUMN].between(TEST_START, TEST_END).to_numpy()

if not train_mask.any() or not test_mask.any():
    raise ValueError("O arquivo deve conter tanto 2022-2023 quanto 2024.")

q_original = pd.to_numeric(df[FLOW_COLUMN], errors="coerce")
non_numeric_flow = df[FLOW_COLUMN].notna() & q_original.isna()
if non_numeric_flow.any():
    raise ValueError("A coluna de vazao contem textos ou valores nao numericos.")

x_all = df[EXOG_COLUMNS].apply(pd.to_numeric, errors="coerce")
if x_all.isna().any().any():
    missing_exog = x_all.isna().sum()
    missing_exog = missing_exog[missing_exog > 0]
    raise ValueError(
        "O Kalman nao preenchera covariaveis exogenas. Existem NaNs em:\n"
        + missing_exog.to_string()
    )

if not np.isfinite(x_all.to_numpy()).all():
    raise ValueError("Existem valores infinitos nas covariaveis exogenas.")

q_model_all = transform_flow(q_original)
train_index = pd.date_range(TRAIN_START, TRAIN_END, freq="h")
test_index = pd.date_range(TEST_START, TEST_END, freq="h")

if not pd.DatetimeIndex(df.loc[train_mask, DATETIME_COLUMN]).equals(train_index):
    raise ValueError("O bloco 2022-2023 nao coincide com a grade horaria esperada.")
if not pd.DatetimeIndex(df.loc[test_mask, DATETIME_COLUMN]).equals(test_index):
    raise ValueError("O bloco 2024 nao coincide com a grade horaria esperada.")

q_train_model = pd.Series(
    q_model_all.loc[train_mask].to_numpy(),
    index=train_index,
    name="log_vazao" if USE_LOG_FLOW else "vazao",
)
q_test_model = pd.Series(
    q_model_all.loc[test_mask].to_numpy(),
    index=test_index,
    name="log_vazao" if USE_LOG_FLOW else "vazao",
)
q_train_original = pd.Series(
    q_original.loc[train_mask].to_numpy(), index=q_train_model.index
)

# Padronizacao exclusivamente com 2022-2023.
x_train_raw = x_all.loc[train_mask].copy()
x_test_raw = x_all.loc[test_mask].copy()
x_mean = x_train_raw.mean(axis=0)
x_std = x_train_raw.std(axis=0, ddof=0)

constant_columns = x_std.index[x_std <= np.finfo(float).eps].tolist()
if constant_columns:
    print("Colunas constantes removidas do ajuste:", constant_columns)

active_exog = [column for column in EXOG_COLUMNS if column not in constant_columns]
if not active_exog:
    raise ValueError("Nenhuma covariavel exogena variavel permaneceu.")

x_train = (x_train_raw[active_exog] - x_mean[active_exog]) / x_std[active_exog]
x_test = (x_test_raw[active_exog] - x_mean[active_exog]) / x_std[active_exog]
x_train.index = q_train_model.index
x_test.index = q_test_model.index

exog_rank = int(np.linalg.matrix_rank(x_train.to_numpy()))
if exog_rank < len(active_exog):
    raise ValueError(
        "A matriz exogena nao possui posto completo "
        f"({exog_rank}/{len(active_exog)}). Revise as dummies antes do ajuste."
    )

condition_number = float(np.linalg.cond(x_train.to_numpy()))
if condition_number > 1e8:
    warnings.warn(
        f"Matriz exogena mal condicionada (condicao={condition_number:.3e}). "
        "Os erros-padrao dos betas podem ser instaveis."
    )

actual_missing = q_original.isna().to_numpy()
missing_train = actual_missing & train_mask
missing_test = actual_missing & test_mask

if actual_missing.sum() == 0:
    raise ValueError("Nao existem NaNs na coluna de vazao; nada seria reconstruido.")

if actual_missing[:2].any():
    raise ValueError(
        "Existe vazao ausente nas duas primeiras horas da serie. Como os lags "
        "de chuva de 1 e 2 h exigiriam precipitacao anterior a 2022, forneca "
        "esses antecedentes antes de executar o Kalman."
    )

if STRICT_INPUT_AUDIT:
    audit_observed = {
        "SHA-256": input_sha256,
        "linhas": int(len(df)),
        "NaNs_totais": int(actual_missing.sum()),
        "NaNs_2022_2023": int(missing_train.sum()),
        "NaNs_2024": int(missing_test.sum()),
    }
    audit_expected = {
        "SHA-256": EXPECTED_INPUT_SHA256,
        "linhas": EXPECTED_ROWS,
        "NaNs_totais": EXPECTED_MISSING_TOTAL,
        "NaNs_2022_2023": EXPECTED_MISSING_TRAIN,
        "NaNs_2024": EXPECTED_MISSING_TEST,
    }
    mismatches = {
        key: {"observado": audit_observed[key], "esperado": audit_expected[key]}
        for key in audit_expected
        if audit_observed[key] != audit_expected[key]
    }
    if mismatches:
        raise RuntimeError(
            "A auditoria congelada do arquivo de entrada falhou:\n"
            + json.dumps(mismatches, ensure_ascii=False, indent=2)
        )

upper_bound = (
    float(FLOW_UPPER_BOUND)
    if FLOW_UPPER_BOUND is not None
    else float(q_original.max(skipna=True))
)
if upper_bound <= FLOW_LOWER_BOUND:
    raise ValueError("O limite superior deve ser maior que 600 L/s.")

print("=" * 78)
print("DADOS VALIDADOS")
print("Arquivo:", INPUT_FILE)
print("SHA-256:", input_sha256)
print("Linhas:", len(df))
print("NaNs totais de vazao:", int(actual_missing.sum()))
print("NaNs em 2022-2023:", int(missing_train.sum()))
print("NaNs em 2024:", int(missing_test.sum()))
print("Exogenas ativas:", len(active_exog))
print("Posto da matriz exogena:", f"{exog_rank}/{len(active_exog)}")
print("Covariaveis de chuva criadas:", ENGINEERED_EXOG_COLUMNS)
print("Limites:", FLOW_LOWER_BOUND, "a", upper_bound, "L/s")
print("Memoria pluviometrica de 20 dias: EXCLUIDA")
print("=" * 78)


# =============================================================================
# 5. AJUSTE DEFINITIVO EM 2022-2023
# =============================================================================

(
    selected_model_spec,
    model_train,
    result_train,
    optimization_attempts_df,
    model_selection_df,
) = fit_and_select_model(q_train_model, x_train)

retvals = getattr(result_train, "mle_retvals", {}) or {}
projected_gradient_norm = optimizer_projected_gradient_norm(result_train)
final_covariance = covariance_diagnostics(result_train)
final_covariance_raw_condition = final_covariance["raw_condition_number"]
final_covariance_correlation_condition = final_covariance[
    "correlation_condition_number"
]
converged = bool(retvals.get("converged", False))

if not (
    converged
    and retvals.get("warnflag", np.nan) == 0
    and projected_gradient_norm <= GRADIENT_ACCEPTANCE_TOLERANCE
    and final_covariance["positive_finite_diagonal"]
    and final_covariance_correlation_condition
    <= COVARIANCE_CORRELATION_CONDITION_MAX
    and not near_zero_estimated_variances(result_train)
):
    raise RuntimeError(
        "O modelo selecionado deixou de satisfazer os criterios numericos "
        "apos a selecao. Nao use esta execucao."
    )

print("Modelo selecionado:", selected_model_spec["model_id"])
print(model_selection_df.to_string(index=False))
print(result_train.summary())


# =============================================================================
# 6. RECONSTRUCOES: SMOOTHER 2022-2023 E FILTRO CAUSAL EM 2024
# =============================================================================

prediction_train_smoothed = result_train.get_prediction(
    start=0,
    end=len(q_train_model) - 1,
    information_set="smoothed",
    signal_only=False,
)
train_converted = prediction_to_original_scale(
    prediction_train_smoothed, FLOW_LOWER_BOUND, upper_bound
)

# extend aplica somente filtragem aos novos dados e preserva os parametros
# estimados em 2022-2023. A previsao "predicted" em t usa dados ate t-1.
result_test = result_train.extend(endog=q_test_model, exog=x_test)
prediction_test_causal = result_test.get_prediction(
    start=0,
    end=len(q_test_model) - 1,
    information_set="predicted",
    signal_only=False,
)
test_converted = prediction_to_original_scale(
    prediction_test_causal, FLOW_LOWER_BOUND, upper_bound
)

estimate_model_scale = np.full(len(df), np.nan)
estimate_raw = np.full(len(df), np.nan)
lower_raw = np.full(len(df), np.nan)
upper_raw = np.full(len(df), np.nan)
estimate_bounded = np.full(len(df), np.nan)
lower_bounded = np.full(len(df), np.nan)
upper_bounded = np.full(len(df), np.nan)
clipped_low = np.zeros(len(df), dtype=bool)
clipped_high = np.zeros(len(df), dtype=bool)

for destination, train_values, test_values in [
    (estimate_model_scale, train_converted["model_scale"], test_converted["model_scale"]),
    (estimate_raw, train_converted["estimate_raw"], test_converted["estimate_raw"]),
    (lower_raw, train_converted["lower_raw"], test_converted["lower_raw"]),
    (upper_raw, train_converted["upper_raw"], test_converted["upper_raw"]),
    (estimate_bounded, train_converted["estimate"], test_converted["estimate"]),
    (lower_bounded, train_converted["lower"], test_converted["lower"]),
    (upper_bounded, train_converted["upper"], test_converted["upper"]),
]:
    destination[train_mask] = train_values
    destination[test_mask] = test_values

clipped_low[train_mask] = train_converted["clipped_low"]
clipped_low[test_mask] = test_converted["clipped_low"]
clipped_high[train_mask] = train_converted["clipped_high"]
clipped_high[test_mask] = test_converted["clipped_high"]

q_final = q_original.to_numpy(dtype=float, copy=True)
q_final[actual_missing] = estimate_bounded[actual_missing]

if not np.isfinite(q_final).all():
    raise RuntimeError("Ainda existem vazoes ausentes ou infinitas apos o Kalman.")

method_by_row = np.full(len(df), "observado_preservado", dtype=object)
method_by_row[missing_train] = "kalman_smoother_2022_2023"
method_by_row[missing_test] = "kalman_filter_causal_2024"

gap_table_df, gap_id_by_row = consecutive_gap_table(
    actual_missing, actual_index, method_by_row
)
gap_table_df["classe_duracao"] = gap_table_df["duracao_horas"].map(
    duration_class
)


# =============================================================================
# 7. VALIDACAO INTERNA POR LACUNAS ARTIFICIAIS
# =============================================================================

if RUN_MASKED_VALIDATION:
    (
        validation_df,
        validation_duration_df,
        validation_blocks_df,
        validation_points_df,
        validation_failures_df,
    ) = masked_validation(
        y_train_original=q_train_original,
        y_train_model=q_train_model,
        x_train=x_train,
        final_params=result_train.params,
        model_spec=selected_model_spec,
        lower_bound=FLOW_LOWER_BOUND,
        upper_bound=upper_bound,
        reference_gap_lengths=gap_table_df["duracao_horas"].tolist(),
    )
else:
    validation_df = pd.DataFrame(
        [{"status": "nao_executada", "motivo": "RUN_MASKED_VALIDATION=False"}]
    )
    validation_duration_df = pd.DataFrame()
    validation_blocks_df = pd.DataFrame()
    validation_points_df = pd.DataFrame()
    validation_failures_df = pd.DataFrame()


# =============================================================================
# 8. TABELAS DE AUDITORIA
# =============================================================================

output_df = df.copy()
output_df.insert(
    output_df.columns.get_loc(FLOW_COLUMN) + 1,
    f"{FLOW_COLUMN}_original_pre_kalman",
    q_original,
)
output_df[FLOW_COLUMN] = q_final
output_df["era_NaN_pre_kalman"] = actual_missing
output_df["gap_id_kalman"] = gap_id_by_row
output_df["metodo_origem_vazao"] = method_by_row
output_df["estimativa_kalman_escala_modelo"] = estimate_model_scale
output_df["estimativa_kalman_sem_limite_L_s"] = estimate_raw
output_df["IC95_inferior_sem_limite_L_s"] = lower_raw
output_df["IC95_superior_sem_limite_L_s"] = upper_raw
output_df["estimativa_kalman_limitada_L_s"] = estimate_bounded
output_df["IC95_inferior_limitado_L_s"] = lower_bounded
output_df["IC95_superior_limitado_L_s"] = upper_bounded
output_df["aplicou_limite_inferior_600"] = clipped_low & actual_missing
output_df["aplicou_limite_superior"] = clipped_high & actual_missing
output_df["alvo_reconstruido_por_kalman_em_2024"] = missing_test
output_df["usar_na_sensibilidade_2024_somente_observados"] = (
    test_mask & ~missing_test
)

innovation_diagnostics_df, standardized_innovations = (
    standardized_innovation_diagnostics(result_train)
)
standardized_innovations_full = np.full(len(df), np.nan)
standardized_innovations_full[train_mask] = standardized_innovations
output_df["erro_previsao_padronizado_filtro_treino"] = (
    standardized_innovations_full
)
numerical_stability_df = numerical_stability_table(result_train)

clipping_audit_df = pd.DataFrame(
    {
        "gap_id": gap_id_by_row[actual_missing].astype(int),
        "metodo": method_by_row[actual_missing],
        "estimativa_sem_limite_L_s": estimate_raw[actual_missing],
        "estimativa_com_limite_L_s": estimate_bounded[actual_missing],
        "limite_inferior_aplicado": clipped_low[actual_missing],
        "limite_superior_aplicado": clipped_high[actual_missing],
    }
).merge(
    gap_table_df[["gap_id", "duracao_horas", "classe_duracao"]],
    on="gap_id",
    how="left",
)
clipping_summary_df = (
    clipping_audit_df.groupby(["metodo", "classe_duracao"], as_index=False)
    .agg(
        n_reconstruido=("gap_id", "size"),
        limite_inferior_n=("limite_inferior_aplicado", "sum"),
        limite_superior_n=("limite_superior_aplicado", "sum"),
        estimativa_bruta_min_L_s=("estimativa_sem_limite_L_s", "min"),
        estimativa_bruta_max_L_s=("estimativa_sem_limite_L_s", "max"),
    )
)
clipping_summary_df["qualquer_limite_n"] = (
    clipping_summary_df["limite_inferior_n"]
    + clipping_summary_df["limite_superior_n"]
)
clipping_summary_df["qualquer_limite_fracao"] = (
    clipping_summary_df["qualquer_limite_n"]
    / clipping_summary_df["n_reconstruido"]
)

param_ci = np.asarray(result_train.conf_int(alpha=ALPHA_INTERVAL))
parameters_df = pd.DataFrame(
    {
        "parametro": list(result_train.model.param_names),
        "estimativa": np.asarray(result_train.params, dtype=float),
        "erro_padrao": np.asarray(result_train.bse, dtype=float),
        "z": np.asarray(result_train.zvalues, dtype=float),
        "p_valor": np.asarray(result_train.pvalues, dtype=float),
        "IC95_inferior": param_ci[:, 0],
        "IC95_superior": param_ci[:, 1],
    }
)

scaler_df = pd.DataFrame(
    {
        "variavel": EXOG_COLUMNS,
        "ativa_no_modelo": [column in active_exog for column in EXOG_COLUMNS],
        "media_2022_2023": [float(x_mean[column]) for column in EXOG_COLUMNS],
        "desvio_padrao_2022_2023": [float(x_std[column]) for column in EXOG_COLUMNS],
    }
)

feature_engineering_df = pd.DataFrame(
    {
        "variavel": ENGINEERED_EXOG_COLUMNS,
        "definicao": [
            FEATURE_ENGINEERING_DEFINITIONS[column]
            for column in ENGINEERED_EXOG_COLUMNS
        ],
        "preexistia_e_foi_verificada": [
            column in preexisting_engineered_columns
            for column in ENGINEERED_EXOG_COLUMNS
        ],
        "criada_nesta_execucao": [
            column not in preexisting_engineered_columns
            for column in ENGINEERED_EXOG_COLUMNS
        ],
    }
)

redundancy_df = pd.DataFrame(
    {
        "variavel_excluida": EXCLUDED_EXACTLY_REDUNDANT_COLUMNS,
        "motivo": [
            (
                "Igual a soma das dummies de segunda a sexta; seria "
                "colinear com as dummies de dia da semana"
            ),
            (
                "Determinada exatamente pela dummy Feriado menos as dummies "
                "Feriado_Manha, Feriado_Noite e Feriado_Tarde"
            ),
        ],
    }
)

# Garante que as matrizes exportadas correspondam exatamente aos parametros da
# solucao selecionada, independentemente da ordem das tentativas multistart.
result_train.model.update(np.asarray(result_train.params, dtype=float))
state_space = result_train.model.ssm
matrix_rows = []
for matrix_name in [
    "design",
    "transition",
    "selection",
    "state_cov",
    "obs_cov",
]:
    matrix_rows.extend(matrix_to_long(matrix_name, state_space[matrix_name]))

filter_results = result_train.filter_results
for matrix_name, matrix_value in [
    ("initial_state", getattr(filter_results, "initial_state", np.array([]))),
    (
        "initial_state_cov",
        getattr(filter_results, "initial_state_cov", np.array([])),
    ),
    (
        "initial_diffuse_state_cov",
        getattr(filter_results, "initial_diffuse_state_cov", np.array([])),
    ),
]:
    if matrix_value is not None and np.asarray(matrix_value).size:
        matrix_rows.extend(matrix_to_long(matrix_name, matrix_value))

matrices_df = pd.DataFrame(matrix_rows)

software_versions = {
    "python": sys.version.replace("\n", " "),
    "platform": platform.platform(),
    "numpy": np.__version__,
    "pandas": pd.__version__,
    "scipy": scipy.__version__,
    "statsmodels": statsmodels.__version__,
    "openpyxl": openpyxl.__version__,
}

metadata = {
    "execution_utc": datetime.now(timezone.utc).isoformat(),
    "input_file": str(INPUT_FILE),
    "input_sheet": INPUT_SHEET,
    "input_sha256": input_sha256,
    "datetime_column": DATETIME_COLUMN,
    "flow_column": FLOW_COLUMN,
    "flow_unit": "L/s",
    "number_rows": int(len(df)),
    "period_start": actual_index.min().isoformat(),
    "period_end": actual_index.max().isoformat(),
    "training_start": TRAIN_START.isoformat(),
    "training_end": TRAIN_END.isoformat(),
    "test_start": TEST_START.isoformat(),
    "test_end": TEST_END.isoformat(),
    "missing_total": int(actual_missing.sum()),
    "missing_2022_2023": int(missing_train.sum()),
    "missing_2024": int(missing_test.sum()),
    "max_gap_hours": int(gap_table_df["duracao_horas"].max()),
    "max_gap_hours_2022_2023": int(
        gap_table_df.loc[
            gap_table_df["metodo"] == "kalman_smoother_2022_2023",
            "duracao_horas",
        ].max()
    ),
    "max_gap_hours_2024": int(
        gap_table_df.loc[
            gap_table_df["metodo"] == "kalman_filter_causal_2024",
            "duracao_horas",
        ].max()
    ),
    "strict_input_audit": STRICT_INPUT_AUDIT,
    "expected_input_sha256": EXPECTED_INPUT_SHA256,
    "engineered_exog_columns": ENGINEERED_EXOG_COLUMNS,
    "engineered_exog_definitions": FEATURE_ENGINEERING_DEFINITIONS,
    "engineered_columns_preexisting_and_verified": preexisting_engineered_columns,
    "engineered_initial_edge_policy": (
        "lag_1h and lag_2h use zero only before the first available timestamp; "
        "the run is rejected if flow is missing in either of the first two rows"
    ),
    "excluded_long_memory_rainfall_features": [
        "precip_acum_20d",
        "interacao_chuva_saturacao",
    ],
    "long_memory_exclusion_reason": (
        "20-day rainfall memory was not retained because it could propagate "
        "antecedent wetness and overestimate reconstructed flow"
    ),
    "excluded_exactly_redundant_columns": EXCLUDED_EXACTLY_REDUNDANT_COLUMNS,
    "exog_requested": EXOG_COLUMNS,
    "exog_active": active_exog,
    "exog_constant_removed": constant_columns,
    "exog_matrix_rank": exog_rank,
    "exog_matrix_number_columns": len(active_exog),
    "exog_condition_number": condition_number,
    "exog_standardization": "media e desvio-padrao estimados somente em 2022-2023",
    "response_transform": "log(Q)" if USE_LOG_FLOW else "Q sem transformacao",
    "selected_model_id": selected_model_spec["model_id"],
    "candidate_model_specs": CANDIDATE_MODELS,
    "candidate_selection_rule": (
        "lowest BIC among candidates satisfying native convergence, projected "
        "gradient, covariance conditioning, and multistart replication"
    ),
    "candidate_selection_table": json_safe(model_selection_df),
    "observation_equation": observation_equation_for_spec(selected_model_spec),
    "state_equations": state_equations_for_spec(selected_model_spec),
    "model_class": "statsmodels.tsa.statespace.structural.UnobservedComponents",
    "state_names": list(result_train.model.state_names),
    "state_dimension": int(result_train.model.k_states),
    "stochastic_level": True,
    "trend_included": bool(selected_model_spec["trend"]),
    "stochastic_trend": bool(selected_model_spec["stochastic_trend"]),
    "autoregressive_order": int(selected_model_spec.get("ar_order", 0)),
    "irregular_observation_error": bool(selected_model_spec["irregular"]),
    "observation_covariance_R_status": (
        "estimated by maximum likelihood"
        if selected_model_spec["irregular"]
        else "fixed at zero by candidate specification"
    ),
    "mle_regression": True,
    "initialization": (
        "exact diffuse for nonstationary states; stationary AR initialization "
        "handled by statsmodels"
        if int(selected_model_spec.get("ar_order", 0)) > 0
        else "exact diffuse for the nonstationary local-level state"
    ),
    "initial_state": json_safe(
        getattr(filter_results, "initial_state", np.array([]))
    ),
    "initial_state_covariance": json_safe(
        getattr(filter_results, "initial_state_cov", np.array([]))
    ),
    "initial_diffuse_state_covariance": json_safe(
        getattr(filter_results, "initial_diffuse_state_cov", np.array([]))
    ),
    "process_covariance_Q_or_equivalent": json_safe(
        np.asarray(state_space["state_cov"])
    ),
    "observation_covariance_R_or_equivalent": json_safe(
        np.asarray(state_space["obs_cov"])
    ),
    "use_exact_diffuse": USE_EXACT_DIFFUSE,
    "parameter_estimation": "maximum likelihood evaluated by Kalman filter",
    "optimizer": OPTIMIZER,
    "optimizer_score_method": (
        "finite differences internal to L-BFGS-B (optim_score=None)"
    ),
    "lbfgs_pgtol": LBFGS_PGTOL,
    "lbfgs_factr": LBFGS_FACTR,
    "lbfgs_memory": LBFGS_MEMORY,
    "lbfgs_maxfun": LBFGS_MAXFUN,
    "lbfgs_finite_difference_epsilon": LBFGS_FINITE_DIFFERENCE_EPSILON,
    "maxiter_final": MAXITER_FINAL,
    "multi_start_attempts_per_candidate": N_STARTS_PER_CANDIDATE,
    "minimum_stable_starts": MIN_STABLE_STARTS,
    "random_seed": RANDOM_SEED,
    "convergence_acceptance_rule": (
        "converged=True; warnflag=0; finite parameters and likelihood; "
        f"projected gradient <= {GRADIENT_ACCEPTANCE_TOLERANCE}; positive "
        "finite covariance diagonal; parameter-correlation condition number "
        f"<= {COVARIANCE_CORRELATION_CONDITION_MAX}; no estimated variance "
        "at the zero boundary; the highest-likelihood numerically valid basin "
        "must meet those criteria; and at least "
        f"{MIN_STABLE_STARTS} starts reproducing the best solution within "
        f"DeltaLLF <= {LOGLIK_CLUSTER_TOLERANCE} and parameter equivalence "
        "according to abs(delta) <= "
        f"{PARAMETER_CLUSTER_ABSOLUTE_TOLERANCE} + "
        f"{PARAMETER_CLUSTER_RELATIVE_TOLERANCE} * "
        "max(abs(parameter), abs(reference))"
    ),
    "converged": converged,
    "projected_gradient_inf_norm": projected_gradient_norm,
    "covariance_raw_condition_number": final_covariance_raw_condition,
    "covariance_correlation_condition_number": (
        final_covariance_correlation_condition
    ),
    "covariance_positive_finite_diagonal": final_covariance[
        "positive_finite_diagonal"
    ],
    "mle_retvals": json_safe(retvals),
    "log_likelihood": float(result_train.llf),
    "aic": float(result_train.aic),
    "bic": float(result_train.bic),
    "hqic": float(result_train.hqic),
    "lower_operational_bound_L_s": FLOW_LOWER_BOUND,
    "upper_operational_bound_L_s": upper_bound,
    "upper_bound_rule": "maximum originally observed valid flow" if FLOW_UPPER_BOUND is None else "manually fixed",
    "back_transform_point_estimate": (
        "exp(predicted log-flow mean), interpreted as conditional median on "
        "the original scale"
    ),
    "clipped_lower_missing_count": int(np.sum(clipped_low & actual_missing)),
    "clipped_upper_missing_count": int(np.sum(clipped_high & actual_missing)),
    "clipped_2022_2023_count": int(
        np.sum((clipped_low | clipped_high) & missing_train)
    ),
    "clipped_2024_count": int(
        np.sum((clipped_low | clipped_high) & missing_test)
    ),
    "clipped_2022_2023_fraction": float(
        np.sum((clipped_low | clipped_high) & missing_train)
        / np.sum(missing_train)
    ),
    "clipped_2024_fraction": float(
        np.sum((clipped_low | clipped_high) & missing_test)
        / np.sum(missing_test)
    ),
    "development_gap_method": "fixed-interval Kalman smoother; conditions on all 2022-2023 observations",
    "test_gap_method": (
        "one-step-ahead Kalman prediction; response observations through t-1 "
        "and exogenous information available at reconstructed hour t"
    ),
    "observed_values_preserved": True,
    "gap_position_counts": json_safe(gap_table_df["posicao"].value_counts()),
    "series_edge_gaps_present": bool(
        gap_table_df["posicao"].isin(
            ["borda_inicial_serie", "borda_final_serie"]
        ).any()
    ),
    "masked_validation_run": RUN_MASKED_VALIDATION,
    "masked_validation_required_successes": (
        VALIDATION_REQUIRED_SUCCESSES if RUN_MASKED_VALIDATION else 0
    ),
    "masked_validation_successes": (
        int(validation_df["repeticao"].nunique())
        if RUN_MASKED_VALIDATION
        else 0
    ),
    "masked_validation_attempt_failures": int(len(validation_failures_df)),
    "masked_validation_max_attempts": VALIDATION_MAX_ATTEMPTS,
    "masked_validation_fraction": VALIDATION_FRACTION if RUN_MASKED_VALIDATION else 0,
    "masked_validation_refits_parameters_after_masking": True,
    "masked_validation_gap_lengths": (
        "sampled from the empirical lengths of all actual 2022-2024 gaps; "
        "artificial blocks are placed only on observed 2022-2023 responses"
    ),
    "masked_validation_limits_sensitivity": ["sem_limites", "com_limites"],
    "masked_validation_duration_classes": VALIDATION_DURATION_CLASSES,
    "bias_definition": "estimate minus observation",
    "innovation_diagnostics": json_safe(innovation_diagnostics_df),
    "numerical_stability_diagnostics": json_safe(numerical_stability_df),
    "software_versions": software_versions,
}

config_df = pd.DataFrame(
    [{"campo": key, "valor": json.dumps(json_safe(value), ensure_ascii=False) if isinstance(value, (list, dict)) else value}
     for key, value in metadata.items()]
)

validation_summary_df = (
    validation_df.groupby(["metodo", "aplicacao_limites"], as_index=False)
    .agg(
        repeticoes=("repeticao", "count"),
        n_total=("n", "sum"),
        MAE_medio_L_s=("MAE_L_s", "mean"),
        RMSE_medio_L_s=("RMSE_L_s", "mean"),
        bias_medio_L_s=("bias_L_s", "mean"),
        NSE_medio=("NSE", "mean"),
        cobertura_95_media=("coverage_95", "mean"),
        largura_IC95_media_L_s=("mean_interval_width_L_s", "mean"),
    )
    if RUN_MASKED_VALIDATION
    else pd.DataFrame()
)

validation_duration_summary_df = (
    validation_duration_df.groupby(
        ["metodo", "aplicacao_limites", "classe_duracao"], as_index=False
    )
    .agg(
        repeticoes=("repeticao", "count"),
        blocos_totais=("numero_blocos", "sum"),
        n_total=("n", "sum"),
        MAE_medio_L_s=("MAE_L_s", "mean"),
        RMSE_medio_L_s=("RMSE_L_s", "mean"),
        bias_medio_L_s=("bias_L_s", "mean"),
        NSE_medio=("NSE", "mean"),
        cobertura_95_media=("coverage_95", "mean"),
        largura_IC95_media_L_s=("mean_interval_width_L_s", "mean"),
    )
    if RUN_MASKED_VALIDATION and not validation_duration_df.empty
    else pd.DataFrame()
)


# =============================================================================
# 9. SALVAMENTO
# =============================================================================

with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
    output_df.to_excel(writer, sheet_name="serie_final", index=False)
    parameters_df.to_excel(writer, sheet_name="parametros", index=False)
    config_df.to_excel(writer, sheet_name="configuracao", index=False)
    gap_table_df.to_excel(writer, sheet_name="lacunas", index=False)
    model_selection_df.to_excel(
        writer, sheet_name="selecao_modelo", index=False
    )
    optimization_attempts_df.to_excel(
        writer, sheet_name="otimizacao", index=False
    )
    scaler_df.to_excel(writer, sheet_name="escala_exogenas", index=False)
    feature_engineering_df.to_excel(
        writer, sheet_name="engenharia_exogenas", index=False
    )
    redundancy_df.to_excel(
        writer, sheet_name="exogenas_redundantes", index=False
    )
    matrices_df.to_excel(writer, sheet_name="matrizes_estado", index=False)
    validation_df.to_excel(
        writer, sheet_name="validacao_mascarada", index=False
    )
    validation_summary_df.to_excel(
        writer, sheet_name="resumo_validacao", index=False
    )
    validation_duration_df.to_excel(
        writer, sheet_name="validacao_duracao", index=False
    )
    validation_duration_summary_df.to_excel(
        writer, sheet_name="resumo_por_duracao", index=False
    )
    validation_blocks_df.to_excel(
        writer, sheet_name="validacao_por_bloco", index=False
    )
    validation_points_df.to_excel(
        writer, sheet_name="validacao_pontos", index=False
    )
    validation_failures_df.to_excel(
        writer, sheet_name="validacao_falhas", index=False
    )
    clipping_summary_df.to_excel(
        writer, sheet_name="limites_resumo", index=False
    )
    clipping_audit_df.to_excel(
        writer, sheet_name="limites_detalhe", index=False
    )
    innovation_diagnostics_df.to_excel(
        writer, sheet_name="diagnosticos_inovacoes", index=False
    )
    numerical_stability_df.to_excel(
        writer, sheet_name="estabilidade_numerica", index=False
    )

format_excel(OUTPUT_XLSX)

metadata["output_xlsx"] = str(OUTPUT_XLSX)
metadata["output_xlsx_sha256"] = sha256_file(OUTPUT_XLSX)

code_copy_saved = False
if "__file__" in globals():
    executed_code_path = Path(__file__).resolve()
    if executed_code_path.exists() and executed_code_path.is_file():
        if executed_code_path != OUTPUT_CODE_COPY.resolve():
            shutil.copy2(executed_code_path, OUTPUT_CODE_COPY)
        code_copy_saved = True
metadata["executed_code_copy_saved"] = code_copy_saved
if code_copy_saved:
    metadata["executed_code_copy"] = str(OUTPUT_CODE_COPY)
    metadata["executed_code_sha256"] = sha256_file(OUTPUT_CODE_COPY)

with open(OUTPUT_JSON, "w", encoding="utf-8") as handle:
    json.dump(json_safe(metadata), handle, ensure_ascii=False, indent=2)

with open(OUTPUT_REQUIREMENTS, "w", encoding="utf-8") as handle:
    for package_name, version in software_versions.items():
        if package_name in {"python", "platform"}:
            continue
        handle.write(f"{package_name}=={version}\n")

with open(OUTPUT_SUMMARY, "w", encoding="utf-8") as handle:
    handle.write(f"SELECTED MODEL: {selected_model_spec['model_id']}\n")
    handle.write(
        f"PROJECTED GRADIENT INF NORM: {projected_gradient_norm}\n"
    )
    handle.write(
        "RAW OPG COVARIANCE CONDITION NUMBER (AUDIT ONLY): "
        f"{final_covariance_raw_condition}\n"
    )
    handle.write(
        "PARAMETER-CORRELATION CONDITION NUMBER (ACCEPTANCE): "
        f"{final_covariance_correlation_condition}\n\n"
    )
    handle.write(result_train.summary().as_text())
    handle.write("\n\nMATRIZES Q E R\n")
    handle.write("state_cov (Q ou equivalente):\n")
    handle.write(np.array2string(np.asarray(state_space["state_cov"])))
    handle.write("\nobs_cov (R ou equivalente):\n")
    handle.write(np.array2string(np.asarray(state_space["obs_cov"])))

if SAVE_MODEL_PICKLE:
    result_train.save(OUTPUT_MODEL, remove_data=False)

write_supplementary_report(
    OUTPUT_SUPPLEMENT,
    metadata,
    parameters_df,
    validation_summary_df,
    validation_duration_summary_df,
    gap_table_df,
)

deliverables = [
    OUTPUT_XLSX,
    OUTPUT_JSON,
    OUTPUT_REQUIREMENTS,
    OUTPUT_SUMMARY,
    OUTPUT_SUPPLEMENT,
]
if SAVE_MODEL_PICKLE:
    deliverables.append(OUTPUT_MODEL)
if code_copy_saved:
    deliverables.append(OUTPUT_CODE_COPY)

manifest = {
    "created_utc": datetime.now(timezone.utc).isoformat(),
    "input": {
        "path": str(INPUT_FILE),
        "sha256": input_sha256,
    },
    "outputs": [
        {
            "file": path.name,
            "bytes": int(path.stat().st_size),
            "sha256": sha256_file(path),
        }
        for path in deliverables
    ],
    "scope_note": (
        "Kalman reconstruction only; Q50 validation is a separate subsequent step"
    ),
}
with open(OUTPUT_MANIFEST, "w", encoding="utf-8") as handle:
    json.dump(json_safe(manifest), handle, ensure_ascii=False, indent=2)

with zipfile.ZipFile(OUTPUT_PACKAGE, "w", compression=zipfile.ZIP_DEFLATED) as archive:
    for path in [*deliverables, OUTPUT_MANIFEST]:
        archive.write(path, arcname=path.name)

print("\n" + "=" * 78)
print("KALMAN CONCLUIDO E SALVO")
print("Excel:", OUTPUT_XLSX)
print("Metadados:", OUTPUT_JSON)
print("Versoes fixadas:", OUTPUT_REQUIREMENTS)
print("Resumo do modelo:", OUTPUT_SUMMARY)
print("Texto-base do suplemento:", OUTPUT_SUPPLEMENT)
print("Manifesto de hashes:", OUTPUT_MANIFEST)
print("Pacote completo:", OUTPUT_PACKAGE)
if code_copy_saved:
    print("Copia exata do codigo executado:", OUTPUT_CODE_COPY)
if SAVE_MODEL_PICKLE:
    print("Objeto statsmodels:", OUTPUT_MODEL)
print("SHA-256 do Excel:", metadata["output_xlsx_sha256"])
print("Valores preservados:", int((~actual_missing).sum()))
print("Valores reconstruidos:", int(actual_missing.sum()))
print("  - smoother 2022-2023:", int(missing_train.sum()))
print("  - filtro causal 2024:", int(missing_test.sum()))
print("Limite inferior aplicado:", int(np.sum(clipped_low & actual_missing)))
print("Limite superior aplicado:", int(np.sum(clipped_high & actual_missing)))
print("Maior lacuna consecutiva:", int(gap_table_df["duracao_horas"].max()), "h")
print("Modelo selecionado:", selected_model_spec["model_id"])
print("Convergencia:", converged)
print("Norma infinita do gradiente projetado:", projected_gradient_norm)
print(
    "Condicionamento bruto da covariancia OPG (auditoria):",
    final_covariance_raw_condition,
)
print(
    "Condicionamento da correlacao dos parametros (aceitacao):",
    final_covariance_correlation_condition,
)
if RUN_MASKED_VALIDATION:
    print(
        "Repeticoes validas da validacao:",
        int(validation_df["repeticao"].nunique()),
    )
    print("Tentativas rejeitadas na validacao:", len(validation_failures_df))
print("=" * 78)
